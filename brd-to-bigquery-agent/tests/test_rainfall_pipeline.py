from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from src.brd_parser.document_reader import read_document
from src.llm.provider import EchoProvider
from src.rainfall.llm_analysis import analyze_rainfall_brd
from src.rainfall.models import MetadataColumn, SourceColumn
from src.rainfall.processing import (
    extract_pipeline_contract,
    inspect_sources,
    make_run_identity,
    validate_rows,
)
from src.rainfall.runner import DEFAULT_BRD, run_dry_run
from src.rainfall.sql import plan_schema_evolution
from src.sql_generator.render_template import sql_string

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = DEFAULT_BRD.parent


@pytest.fixture
def rainfall_spec(rainfall_brd_available):
    return extract_pipeline_contract(
        rainfall_brd_available,
        run_date=date(2026, 10, 15),
        run_id="b" * 20,
        idempotency_key="c" * 64,
        ingested_at_utc="2026-10-15T00:00:00Z",
        project_id="rainfall-test-123456",
        location="US",
        gcs_bucket="rainfall-test-bucket-123",
        dataset_00="00_rainfall_monitoring",
        dataset_00_is_approved=True,
        llm_provider=None,
        llm_model=None,
        llm_credential_configured=False,
    )[0]


def test_actual_brd_extracts_three_complete_contracts(rainfall_spec):
    by_id = {feed.feed_id: feed for feed in rainfall_spec.feeds}
    assert set(by_id) == {
        "rain_gauge_daily_observations",
        "watershed_station_map",
        "flood_watch_thresholds",
    }
    assert [column.name for column in by_id["watershed_station_map"].source_columns] == [
        "watershed_id",
        "watershed_label",
        "region_code",
        "station_id",
        "station_latitude",
        "station_longitude",
        "station_elevation_m",
        "watershed_area_km2",
        "effective_start_date",
        "effective_end_date",
    ]
    assert by_id["watershed_station_map"].business_key == [
        "watershed_id",
        "station_id",
        "effective_start_date",
    ]
    assert by_id["watershed_station_map"].partition_field == "effective_start_date"
    assert by_id["watershed_station_map"].cluster_fields == ["watershed_id", "station_id"]
    assert [column.name for column in rainfall_spec.metadata_columns] == [
        "source_file_name",
        "ingested_at",
        "run_date",
    ]
    assert not any(
        column.name in {"source_file_name", "ingested_at", "run_date"}
        for feed in rainfall_spec.feeds
        for column in feed.source_columns
    )


def test_contract_rejects_unknown_fields_bad_types_and_nullability():
    with pytest.raises(ValidationError):
        SourceColumn(
            name="amount",
            type="FLOAT",
            nullable=True,
            required_indicator="No",
            description="amount",
            evidence={
                "requirement": "amount",
                "brd_section": "5.2",
                "confidence": 1,
                "deterministic_validation": "PASS",
                "human_review_required": False,
            },
        )
    with pytest.raises(ValidationError):
        SourceColumn(
            name="amount",
            type="FLOAT64",
            nullable=False,
            required_indicator="No",
            description="amount",
            evidence={
                "requirement": "amount",
                "brd_section": "5.2",
                "confidence": 1,
                "deterministic_validation": "PASS",
                "human_review_required": False,
            },
            unrecognized=True,
        )
    with pytest.raises(ValidationError):
        SourceColumn(
            name="amount",
            type="FLOAT64",
            nullable=True,
            required_indicator="Yes",
            description="amount",
            evidence={
                "requirement": "amount",
                "brd_section": "5.2",
                "confidence": 1,
                "deterministic_validation": "PASS",
                "human_review_required": False,
            },
        )


def test_contract_rejects_unsupported_partition_cluster_and_metadata_type(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    payload = feed.model_dump()
    payload["cluster_fields"] = ["rainfall_mm"]
    with pytest.raises(ValidationError):
        type(feed).model_validate(payload)

    payload = feed.model_dump()
    payload["partition_field"] = "rainfall_mm"
    with pytest.raises(ValidationError):
        type(feed).model_validate(payload)

    with pytest.raises(ValidationError):
        MetadataColumn(
            name="run_date",
            type="STRING",
            description="wrong type",
            evidence={
                "requirement": "run_date",
                "brd_section": "5.5",
                "confidence": 1,
                "deterministic_validation": "PASS",
                "human_review_required": False,
            },
        )


def test_source_contract_validation_rejects_negative_rainfall_and_bad_quality(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    row = {
        "station_id": "SYN-GAUGE-001",
        "observation_date": "2026-10-01",
        "rainfall_mm": "-1",
        "rainfall_duration_hours": "3",
        "peak_intensity_mm_per_hour": "1",
        "quality_flag": "UNKNOWN",
        "source_updated_at": "2026-10-02T00:15:00Z",
    }
    issues, accepted, errors = validate_rows(
        feed.feed_id, feed, [row], "rainfall.csv"
    )
    reasons = {issue.reason for issue in issues}
    assert not accepted
    assert 2 in errors
    assert any("non-negative" in reason for reason in reasons)
    assert any("VALID, ESTIMATED, or MISSING" in reason for reason in reasons)


def test_source_contract_validation_enforces_missing_quality_and_duration(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    row = {
        "station_id": "SYN-GAUGE-002",
        "observation_date": "2026-10-01",
        "rainfall_mm": "0",
        "rainfall_duration_hours": "1",
        "peak_intensity_mm_per_hour": "0",
        "quality_flag": "MISSING",
        "source_updated_at": "",
    }
    issues, accepted, _ = validate_rows(feed.feed_id, feed, [row], "missing.csv")
    assert not accepted
    assert any("MISSING quality requires a null measurement" in issue.reason for issue in issues)
    assert any("Duration must be zero" in issue.reason for issue in issues)


def test_required_rainfall_date_cannot_be_blank(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    row = {
        "station_id": "SYN-GAUGE-003",
        "observation_date": "",
        "rainfall_mm": "1",
        "rainfall_duration_hours": "1",
        "peak_intensity_mm_per_hour": "1",
        "quality_flag": "VALID",
        "source_updated_at": "",
    }
    issues, accepted, _ = validate_rows(feed.feed_id, feed, [row], "blank-date.csv")
    assert not accepted
    assert any(issue.field == "observation_date" for issue in issues)


def test_rainfall_business_key_duplicates_and_bad_timestamp_are_rejected(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    row = {
        "station_id": "SYN-GAUGE-004",
        "observation_date": "2026-10-01",
        "rainfall_mm": "1",
        "rainfall_duration_hours": "1",
        "peak_intensity_mm_per_hour": "1",
        "quality_flag": "VALID",
        "source_updated_at": "2026-10-01 01:00:00",
    }
    issues, accepted, _ = validate_rows(feed.feed_id, feed, [row, row], "duplicate.csv")
    assert not accepted
    assert sum("Duplicate business key" in issue.reason for issue in issues) == 2
    assert sum("include a timezone" in issue.reason for issue in issues) == 2


def test_source_contract_validation_rejects_invalid_effective_date(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "watershed_station_map")
    row = {
        "watershed_id": "W-01",
        "watershed_label": "Basin",
        "region_code": "R-01",
        "station_id": "S-01",
        "station_latitude": "35.5",
        "station_longitude": "-120.1",
        "station_elevation_m": "80",
        "watershed_area_km2": "10",
        "effective_start_date": "2026-10-01",
        "effective_end_date": "2026-09-30",
    }
    issues, accepted, _ = validate_rows(feed.feed_id, feed, [row], "mapping.csv")
    assert not accepted
    assert any("before effective start date" in issue.reason for issue in issues)


def test_mapping_effective_periods_must_not_overlap(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "watershed_station_map")
    base = {
        "watershed_id": "W-02",
        "watershed_label": "Basin",
        "region_code": "R-01",
        "station_id": "S-02",
        "station_latitude": "35.5",
        "station_longitude": "-120.1",
        "station_elevation_m": "80",
        "watershed_area_km2": "10",
    }
    rows = [
        {
            **base,
            "effective_start_date": "2026-01-01",
            "effective_end_date": "2026-01-31",
        },
        {
            **base,
            "effective_start_date": "2026-01-31",
            "effective_end_date": "",
        },
    ]
    issues, accepted, _ = validate_rows(feed.feed_id, feed, rows, "mapping.csv")
    assert not accepted
    assert sum("Effective periods overlap" in issue.reason for issue in issues) == 2


def test_mapping_coordinate_and_area_ranges_are_enforced(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "watershed_station_map")
    row = {
        "watershed_id": "W-03",
        "watershed_label": "Basin",
        "region_code": "R-01",
        "station_id": "S-03",
        "station_latitude": "91",
        "station_longitude": "-181",
        "station_elevation_m": "80",
        "watershed_area_km2": "0",
        "effective_start_date": "2026-10-01",
        "effective_end_date": "",
    }
    issues, accepted, _ = validate_rows(feed.feed_id, feed, [row], "range.csv")
    assert not accepted
    assert {issue.field for issue in issues} >= {
        "station_latitude",
        "station_longitude",
        "watershed_area_km2",
    }


def test_source_contract_validation_checks_threshold_order(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "flood_watch_thresholds")
    base = {
        "watershed_id": "W-01",
        "rolling_window_days": "1",
        "minimum_station_coverage_pct": "80",
        "valid_from": "2026-10-01",
        "valid_to": "2026-12-31",
    }
    rows = [
        {**base, "rainfall_threshold_mm": "30", "severity_level": "WATCH"},
        {**base, "rainfall_threshold_mm": "20", "severity_level": "WARNING"},
    ]
    issues, accepted, _ = validate_rows(feed.feed_id, feed, rows, "threshold.csv")
    assert not accepted
    assert sum("Higher severity threshold is lower" in issue.reason for issue in issues) == 2


def test_threshold_windows_ranges_and_dates_are_enforced(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "flood_watch_thresholds")
    row = {
        "watershed_id": "W-04",
        "rolling_window_days": "2",
        "rainfall_threshold_mm": "0",
        "minimum_station_coverage_pct": "101",
        "severity_level": "CRITICAL",
        "valid_from": "2026-10-02",
        "valid_to": "2026-10-01",
    }
    issues, accepted, _ = validate_rows(feed.feed_id, feed, [row], "invalid-threshold.csv")
    assert not accepted
    assert {issue.field for issue in issues} >= {
        "rolling_window_days",
        "rainfall_threshold_mm",
        "minimum_station_coverage_pct",
        "severity_level",
        "valid_to",
    }


def test_schema_evolution_adds_only_nullable_fields_and_preserves_history(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    existing = [
        {"name": column.name, "type": column.type, "nullable": column.nullable}
        for column in feed.source_columns
        if column.name != "source_updated_at"
    ]
    existing.extend(
        [
            {"name": "source_file_name", "type": "STRING", "nullable": False},
            {"name": "ingested_at", "type": "TIMESTAMP", "nullable": False},
            {"name": "run_date", "type": "DATE", "nullable": False},
            {"name": "legacy_note", "type": "STRING", "nullable": True},
        ]
    )
    plan = plan_schema_evolution(
        existing,
        feed,
        project_id=rainfall_spec.project_id,
        dataset_id=rainfall_spec.dataset_00,
    )
    assert plan["action"] == "ADD_NULLABLE_COLUMNS"
    assert plan["additions"] == [
        {"name": "source_updated_at", "type": "TIMESTAMP", "nullable": True}
    ]
    assert "legacy_note" in plan["preserved_historical_columns"]
    assert "ADD COLUMN IF NOT EXISTS source_updated_at TIMESTAMP" in plan["alter_statements"][0]


def test_schema_evolution_quarantines_incompatible_types(rainfall_spec):
    feed = next(item for item in rainfall_spec.feeds if item.feed_id == "rain_gauge_daily_observations")
    existing = [
        {"name": column.name, "type": column.type, "nullable": column.nullable}
        for column in feed.source_columns
    ]
    existing[0]["type"] = "INT64"
    plan = plan_schema_evolution(existing, feed)
    assert plan["action"] == "QUARANTINE"
    assert "no type alteration" in plan["reason"]
    assert plan["additions"] == []


def test_sql_string_escapes_quotes_and_control_characters():
    assert sql_string('a"b\\c\nd\re\tf') == 'a\\"b\\\\c\\nd\\re\\tf'


def test_real_source_inventory_and_dry_run_are_complete_without_credentials(
    tmp_path, rainfall_sources_available,
):
    report = run_dry_run(
        brd_path=rainfall_sources_available / DEFAULT_BRD.name,
        source_root=rainfall_sources_available,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ={},
    )
    assert report["status"] == "DRY_RUN_COMPLETE_CONFIGURATION_REQUIRED"
    assert report["cloud_operations_performed"] is False
    assert report["llm"]["credential_configured"] is False
    assert report["sql_files"] == []
    assert report["row_counts"]["rain_gauge_daily_observations"]["accepted_rows"] == 271
    assert report["row_counts"]["watershed_station_map"]["accepted_rows"] == 260
    assert report["row_counts"]["flood_watch_thresholds"]["accepted_rows"] == 252
    inventory = json.loads(Path(report["source_inventory_json"]).read_text(encoding="utf-8"))
    assert len(inventory) == 3
    assert all(item["load_eligible"] for item in inventory)
    assert all(item["rejected_rows"] == 0 for item in inventory)
    manifest = json.loads(Path(report["artifacts"]["run_manifest"]).read_text(encoding="utf-8"))
    assert manifest["brd_file"] == DEFAULT_BRD.name
    assert manifest["template_version"] == "rainfall_00_v1"
    assert "Z" == manifest["ingested_at_utc"][-1]

    repeat = run_dry_run(
        brd_path=rainfall_sources_available / DEFAULT_BRD.name,
        source_root=rainfall_sources_available,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ={},
    )
    assert repeat["run_id"] == report["run_id"]
    assert repeat["ingested_at_utc"] == report["ingested_at_utc"]


def test_configured_dry_run_renders_three_placeholder_free_00_scripts(
    tmp_path, rainfall_sources_available,
):
    environment = {
        "GOOGLE_CLOUD_PROJECT": "rainfall-test-123456",
        "GOOGLE_CLOUD_LOCATION": "US",
        "GCS_LANDING_BUCKET": "rainfall-test-bucket-123",
        "BQ_DATASET_00": "00_rainfall_monitoring",
    }
    report = run_dry_run(
        brd_path=rainfall_sources_available / DEFAULT_BRD.name,
        source_root=rainfall_sources_available,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ=environment,
    )
    assert report["status"] == "DRY_RUN_READY_FOR_APPROVAL"
    assert len(report["sql_files"]) == 3
    for filename in report["sql_files"]:
        sql = Path(filename).read_text(encoding="utf-8")
        assert "CREATE OR REPLACE EXTERNAL TABLE" in sql
        assert "CREATE TABLE IF NOT EXISTS" in sql
        assert "BEGIN TRANSACTION;" in sql
        assert "source_file_name STRING NOT NULL" in sql
        assert "ingested_at TIMESTAMP NOT NULL" in sql
        assert "run_date DATE NOT NULL" in sql
        assert "{{" not in sql and "}}" not in sql
        assert "<approved-" not in sql
        assert "source_file_name" not in sql.split("OPTIONS (", 1)[0].split("EXTERNAL TABLE", 1)[1]


def test_sql_generation_escapes_untrusted_source_filename(tmp_path, rainfall_brd_available):
    source_root = tmp_path / "sources"
    source_root.mkdir()
    source = source_root / 'station "daily.csv'
    source.write_text(
        "station_id,observation_date,rainfall_mm,rainfall_duration_hours,peak_intensity_mm_per_hour,quality_flag,source_updated_at\n"
        "SYN-GAUGE-005,2026-10-01,1,1,1,VALID,2026-10-01T01:00:00Z\n",
        encoding="utf-8",
    )
    report = run_dry_run(
        brd_path=rainfall_brd_available,
        source_root=source_root,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ={
            "GOOGLE_CLOUD_PROJECT": "rainfall-test-123456",
            "GOOGLE_CLOUD_LOCATION": "US",
            "GCS_LANDING_BUCKET": "rainfall-test-bucket-123",
            "BQ_DATASET_00": "00_rainfall_monitoring",
        },
    )
    assert len(report["sql_files"]) == 1
    sql = Path(report["sql_files"][0]).read_text(encoding="utf-8")
    assert r'station \"daily.csv' in sql
    assert "{{" not in sql and "}}" not in sql


def test_configured_llm_credential_is_reported_only_as_a_boolean(
    tmp_path, rainfall_sources_available,
):
    marker = "do-not-serialize-test-marker"
    report = run_dry_run(
        brd_path=rainfall_sources_available / DEFAULT_BRD.name,
        source_root=rainfall_sources_available,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ={
            "LLM_PROVIDER": "openai",
            "LLM_MODEL": "test-model",
            "LLM_API_KEY": marker,
        },
    )
    assert report["llm"]["credential_configured"] is True
    output_text = "\n".join(
        path.read_text(encoding="utf-8")
        for path in (tmp_path / "generated").rglob("*")
        if path.is_file()
    )
    assert marker not in output_text


def test_example_placeholders_do_not_enable_llm_or_cloud_generation(
    tmp_path, rainfall_sources_available,
):
    report = run_dry_run(
        brd_path=rainfall_sources_available / DEFAULT_BRD.name,
        source_root=rainfall_sources_available,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ={
            "LLM_PROVIDER": "<provider-name>",
            "LLM_MODEL": "<approved-model-name>",
            "LLM_API_KEY": "<set-locally-do-not-commit>",
            "GOOGLE_CLOUD_PROJECT": "<approved-project-id>",
            "GOOGLE_CLOUD_LOCATION": "<approved-region>",
            "GCS_LANDING_BUCKET": "<approved-bucket>",
            "BQ_DATASET_00": "<approved-00-dataset>",
        },
    )
    assert report["llm"]["credential_configured"] is False
    assert report["llm"]["status"] == "NOT_CONFIGURED"
    assert report["sql_files"] == []
    assert not report["resource_configuration"]["project_id_configured"]


def test_required_llm_blocks_sql_planning_when_unconfigured(
    tmp_path, rainfall_sources_available,
):
    report = run_dry_run(
        brd_path=rainfall_sources_available / DEFAULT_BRD.name,
        source_root=rainfall_sources_available,
        output_root=tmp_path / "generated",
        run_date=date(2026, 10, 15),
        environ={
            "GOOGLE_CLOUD_PROJECT": "rainfall-test-123456",
            "GOOGLE_CLOUD_LOCATION": "US",
            "GCS_LANDING_BUCKET": "rainfall-test-bucket-123",
            "BQ_DATASET_00": "00_rainfall_monitoring",
        },
        require_llm=True,
    )
    assert report["status"] == "DRY_RUN_BLOCKED_LLM_REVIEW"
    assert report["sql_files"] == []
    assert report["llm"]["status"] == "NOT_CONFIGURED"
    assert report["cloud_operations_performed"] is False
    assert any(
        item.get("decision") == "required_llm_review"
        for item in report["unresolved_decisions"]
    )


def test_llm_analysis_is_structured_evidence_bound_and_advisory(
    tmp_path, monkeypatch, rainfall_brd_available,
):
    document = read_document(rainfall_brd_available)
    evidence = "Measurement quality status; allowed values: VALID, ESTIMATED, MISSING."
    assert evidence in document.text
    response_payload = {
        "summary": "Reviewed the explicit measurement-quality rule.",
        "findings": [
            {
                "category": "conditional_validation",
                "feed_id": "rain_gauge_daily_observations",
                "source_file": None,
                "field_names": ["quality_flag"],
                "statement": "The MISSING status controls whether measurements may be null.",
                "brd_section": "5.2 Feed A — daily gauge observations",
                "evidence_quote": evidence,
                "confidence": 0.98,
                "human_review_required": True,
            }
        ],
    }
    fake_provider = EchoProvider([json.dumps(response_payload)])
    provider_options = {}

    def resolve_provider(_name, _model, **options):
        provider_options.update(options)
        return fake_provider

    monkeypatch.setattr("src.rainfall.runner.get_provider", resolve_provider)
    report = run_dry_run(
        brd_path=rainfall_brd_available,
        source_root=rainfall_brd_available.parent,
        run_date=date(2026, 10, 15),
        output_root=tmp_path / "generated",
        environ={
            "LLM_PROVIDER": "litellm",
            "LLM_MODEL": "openai/unit-test-model",
            "LLM_API_KEY": "test-only-marker",
            "LLM_API_BASE": "https://llm.example.test/v1",
        },
    )
    assert report["llm"]["status"] == "COMPLETE"
    assert report["llm"]["findings"][0]["field_names"] == ["quality_flag"]
    assert len(fake_provider.calls) == 1
    assert provider_options == {
        "api_key": "test-only-marker",
        "api_base": "https://llm.example.test/v1",
    }
    generated_text = "\n".join(
        path.read_text(encoding="utf-8")
        for path in (tmp_path / "generated").rglob("*")
        if path.is_file()
    )
    assert "test-only-marker" not in generated_text
    analysis_path = Path(report["artifacts"]["llm_analysis"])
    assert analysis_path.is_file()
    assert json.loads(analysis_path.read_text(encoding="utf-8"))["status"] == "COMPLETE"
    assert report["cloud_operations_performed"] is False


def test_llm_analysis_rejects_evidence_not_in_brd(rainfall_spec, rainfall_brd_available):
    document = read_document(rainfall_brd_available)
    payload = {
        "summary": "Unsupported claim.",
        "findings": [
            {
                "category": "unresolved_decision",
                "feed_id": None,
                "source_file": None,
                "field_names": [],
                "statement": "A fabricated requirement.",
                "brd_section": "5.2 Feed A — daily gauge observations",
                "evidence_quote": "this exact quotation does not occur in the source BRD",
                "confidence": 1,
                "human_review_required": True,
            }
        ],
    }
    provider = EchoProvider([json.dumps(payload)])
    with pytest.raises(ValueError, match="evidence not present"):
        analyze_rainfall_brd(document, rainfall_spec, [], provider)


def test_llm_reasoning_repairs_invalid_structured_response_once(rainfall_spec, rainfall_brd_available):
    provider = EchoProvider([
        json.dumps({"summary": "Invalid extra field.", "findings": [], "unexpected": True}),
        json.dumps({"summary": "Validated review.", "findings": []}),
    ])
    analysis, _ = analyze_rainfall_brd(
        read_document(rainfall_brd_available), rainfall_spec, [], provider, max_attempts=2
    )
    assert analysis.summary == "Validated review."
    assert len(provider.calls) == 2
    assert "response_json_schema" in provider.calls[0][1]
    assert "failed deterministic validation" in provider.calls[1][1]


def test_local_source_file_is_not_overwritten(tmp_path, rainfall_brd_available):
    source = tmp_path / "regional_rainfall_daily_observations.csv"
    source.write_text(
        "station_id,observation_date,rainfall_mm,rainfall_duration_hours,peak_intensity_mm_per_hour,quality_flag,source_updated_at\n"
        "SYN-GAUGE-001,2026-10-01,1,1,1,VALID,2026-10-01T01:00:00Z\n",
        encoding="utf-8",
    )
    original = source.read_bytes()
    identity, key, _manifest = make_run_identity("a" * 64, date(2026, 10, 15), [source])
    spec = extract_pipeline_contract(
        rainfall_brd_available,
        run_date=date(2026, 10, 15),
        run_id=identity,
        idempotency_key=key,
        ingested_at_utc="2026-10-15T00:00:00Z",
        project_id=None,
        location=None,
        gcs_bucket=None,
        dataset_00=None,
        dataset_00_is_approved=False,
        llm_provider=None,
        llm_model=None,
        llm_credential_configured=False,
    )[0]
    inspect_sources(
        [source],
        spec.feeds,
        tmp_path / "staging",
        identity,
        date(2026, 10, 15),
        "2026-10-15T00:00:00Z",
        None,
    )
    assert source.read_bytes() == original
    staged = list((tmp_path / "staging").rglob("*.csv"))
    assert len(staged) == 1
    assert "source_file_name" not in staged[0].read_text(encoding="utf-8").splitlines()[0]


def test_excel_sheet_is_converted_without_losing_formatted_identifier(tmp_path, rainfall_spec):
    source = tmp_path / "observations.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "daily_observations"
    sheet.append([
        "station_id",
        "observation_date",
        "rainfall_mm",
        "rainfall_duration_hours",
        "peak_intensity_mm_per_hour",
        "quality_flag",
        "source_updated_at",
    ])
    sheet.append([1, "2026-10-01", 1, 1, 1, "VALID", "2026-10-01T01:00:00Z"])
    sheet["A2"].number_format = "000"
    workbook.save(source)
    original = source.read_bytes()

    prepared = inspect_sources(
        [source],
        rainfall_spec.feeds,
        tmp_path / "staging",
        rainfall_spec.run_id,
        rainfall_spec.run_date,
        rainfall_spec.ingested_at_utc,
        None,
    )

    assert len(prepared) == 1
    assert prepared[0].item.sheet == "daily_observations"
    assert prepared[0].item.load_eligible
    assert prepared[0].item.accepted_rows == 1
    assert source.read_bytes() == original
    assert prepared[0].staging_path is not None
    with prepared[0].staging_path.open(encoding="utf-8") as staged_file:
        staged_rows = list(csv.DictReader(staged_file))
    assert staged_rows[0]["station_id"] == "001"
    assert "source_file_name" not in staged_rows[0]


def test_multiple_plausible_excel_sheets_are_quarantined(tmp_path, rainfall_spec):
    source = tmp_path / "ambiguous.xlsx"
    workbook = Workbook()
    headers = [
        "station_id",
        "observation_date",
        "rainfall_mm",
        "rainfall_duration_hours",
        "peak_intensity_mm_per_hour",
        "quality_flag",
        "source_updated_at",
    ]
    row = ["S-01", "2026-10-01", 1, 1, 1, "VALID", "2026-10-01T01:00:00Z"]
    first = workbook.active
    first.title = "sheet_one"
    first.append(headers)
    first.append(row)
    second = workbook.create_sheet("sheet_two")
    second.append(headers)
    second.append(row)
    workbook.save(source)

    prepared = inspect_sources(
        [source],
        rainfall_spec.feeds,
        tmp_path / "staging",
        rainfall_spec.run_id,
        rainfall_spec.run_date,
        rainfall_spec.ingested_at_utc,
        None,
    )

    assert {item.item.sheet for item in prepared} == {"sheet_one", "sheet_two"}
    assert all(not item.item.load_eligible for item in prepared)
    assert all(
        any("Multiple sheets plausibly match" in reason for reason in item.item.quarantine_reasons)
        for item in prepared
    )
