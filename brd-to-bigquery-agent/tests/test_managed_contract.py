from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from src.brd_parser.document_reader import BrdDocument, BrdLine
from src.orchestrator.managed_contract import (
    ContractField, ManagedContract, ManagedTable, RowRule, checked_select, propose_contract, validate_contract,
)
from src.llm.provider import EchoProvider
from src.orchestrator.managed_runner import compile_publication, inspect_registered_source
from src.rainfall.executor import ContractBlocked

RAW = "clgx-surveillance-dev-9e5e.sample_00.00_readings"
CURATED = "clgx-surveillance-dev-9e5e.sample_01.01_readings"


def raw_table(**overrides):
    values = {
        "table_id": RAW, "layer": "00",
        "columns": [
            ContractField(name="station", type="STRING", nullable=False,
                          description="Required station", evidence_quote="station STRING required"),
            ContractField(name="amount", type="FLOAT64", nullable=False,
                          description="Nonnegative amount", evidence_quote="amount FLOAT64 required"),
        ],
        "business_keys": ["station"], "partition_by": None, "cluster_by": ["station"],
        "source_filename": "readings.csv", "gcs_prefix": "readings/",
        "dependencies": [], "evidence_quote": "00_readings is a managed native table.",
        "row_rules": [RowRule(field="amount", operator="minimum", values=["0"],
                              evidence_quote="amount must be nonnegative")],
    }
    values.update(overrides)
    return ManagedTable(**values)


def curated_table():
    return ManagedTable(
        table_id=CURATED, layer="01",
        columns=raw_table().columns, business_keys=["station"], partition_by=None,
        cluster_by=["station"], dependencies=[RAW],
        select_sql=f"SELECT station, amount FROM `{RAW}`",
        evidence_quote="01_readings is a one-to-one typed copy.",
    )


def contract():
    return ManagedContract(
        project="clgx-surveillance-dev-9e5e", location="US", bucket="sample-dev-private",
        tables=[raw_table(), curated_table()], unresolved=[], unsupported=[],
        execution_summary="Typed source and copy.", resource_evidence="Use sample-dev-private bucket.",
    )


def document():
    text = (
        f"clgx-surveillance-dev-9e5e\nUse sample-dev-private bucket.\n{RAW}\n{CURATED}\n"
        "00_readings is a managed native table.\n01_readings is a one-to-one typed copy.\n"
        "station STRING required\namount FLOAT64 required\namount must be nonnegative\n"
        "readings.csv\nreadings/\n"
    )
    return BrdDocument(path=Path("example.txt"), lines=[BrdLine(index=0, text=text)])


@pytest.mark.parametrize("sql", [
    f"DELETE FROM `{RAW}` WHERE TRUE",
    f"SELECT * FROM `{RAW}`; DROP TABLE `{CURATED}`",
    "SELECT * FROM `foreign.dataset.table`",
    "SELECT * FROM EXTERNAL_QUERY('connection', 'SELECT secret')",
    f"SELECT `foreign.dataset.udf`(amount) FROM `{RAW}`",
    f"SELECT custom_udf(amount) FROM `{RAW}`",
    f"SELECT * FROM `{RAW}` WHERE amount > @user_parameter",
    f"SELECT * FROM `{RAW}` WHERE amount > @@query_label",
])
def test_readonly_ast_blocks_mutations_foreign_reads_remote_udfs_and_parameters(sql):
    with pytest.raises(ValueError):
        checked_select(sql, [RAW])


def test_readonly_ast_allows_explicit_local_cte_query():
    sql = f"WITH inputs AS (SELECT station, amount FROM `{RAW}`) SELECT station, SUM(amount) amount FROM inputs GROUP BY station"
    assert "SUM(amount)" in checked_select(sql, [RAW])


def test_generic_contract_does_not_assume_rainfall_names():
    value = contract()
    validate_contract(value, document())
    sql = compile_publication(value)
    assert f"MERGE `{CURATED}`" in sql
    assert "ASSERT NOT EXISTS" in sql
    assert sql.index("ASSERT NOT EXISTS") < sql.index("BEGIN TRANSACTION") < sql.index("MERGE")
    assert "DROP" not in sql and "TRUNCATE" not in sql


def test_generic_02_optional_contract_has_closed_dependencies_and_separate_transaction():
    value = contract()
    date_column = ContractField(
        name="assessment_date", type="DATE", nullable=False,
        description="Assessment date", evidence_quote="assessment_date DATE required",
    )
    value.tables[1].columns.append(date_column)
    value.tables[1].select_sql = f"SELECT station, amount, @run_date AS assessment_date FROM `{RAW}`"
    kpi = ManagedTable(
        table_id=CURATED.replace("sample_01.01_", "sample_02.02_"), layer="02",
        columns=value.tables[1].columns, business_keys=["station", "assessment_date"],
        partition_by="assessment_date", cluster_by=["station"], dependencies=[CURATED],
        select_sql=f"SELECT station, amount, assessment_date FROM `{CURATED}`",
        evidence_quote="02_readings copies the published curated table.",
    )
    value.tables.append(kpi)
    evidence = document()
    evidence.lines[0].text += (
        f"\n{kpi.table_id}\nassessment_date DATE required\n"
        "02_readings copies the published curated table."
    )
    validate_contract(value, evidence)
    assert f"MERGE `{kpi.table_id}`" not in compile_publication(value)
    from src.sql_generator.dsp import compile_dsp_publication
    assert f"MERGE `{kpi.table_id}`" in compile_dsp_publication([kpi])
    kpi.dependencies = [RAW]
    with pytest.raises(ValueError, match="only published 01"):
        validate_contract(value, evidence)


def test_generic_llm_contract_is_compiled_only_after_deterministic_validation():
    provider = EchoProvider(responses=[contract().model_dump_json()])
    result, response = propose_contract(document(), provider)
    assert [table.table_id for table in result.tables] == [RAW, CURATED]
    assert response["provider"] == "echo"
    assert len(provider.calls) == 1


def test_generic_contract_uses_configured_project_and_rejects_mismatch():
    replacement = "sample-project-123"
    alternate_raw = RAW.replace("clgx-surveillance-dev-9e5e", replacement)
    alternate_curated = CURATED.replace("clgx-surveillance-dev-9e5e", replacement)
    value = contract()
    value.project = replacement
    value.tables[0].table_id = alternate_raw
    value.tables[1].table_id = alternate_curated
    value.tables[1].dependencies = [alternate_raw]
    value.tables[1].select_sql = f"SELECT station, amount FROM `{alternate_raw}`"
    evidence = document()
    evidence.lines[0].text = evidence.lines[0].text.replace("clgx-surveillance-dev-9e5e", replacement)

    validate_contract(value, evidence, expected_project=replacement)
    with pytest.raises(ValueError, match="GOOGLE_CLOUD_PROJECT"):
        validate_contract(value, evidence, expected_project="different-project-123")


def test_generic_llm_validation_repair_is_bounded():
    provider = EchoProvider(responses=['{"invented": true}', contract().model_dump_json()])
    result, _ = propose_contract(document(), provider)
    assert result.project == "clgx-surveillance-dev-9e5e"
    assert len(provider.calls) == 2
    assert "repair" in provider.calls[1][1]


def test_sql_syntax_errors_are_actionable_without_echoing_model_output():
    with pytest.raises(ValueError, match="syntax parsing"):
        checked_select("SELECT ( FROM bad", [])


def test_generic_contract_rejects_missing_grounding_and_forward_dependencies():
    value = contract()
    value.tables[1].columns[0].evidence_quote = "A fabricated unsupported contract definition."
    with pytest.raises(ValueError, match="evidence"):
        validate_contract(value, document())
    value = contract()
    value.tables = list(reversed(value.tables))
    with pytest.raises(ValueError, match="dependencies"):
        validate_contract(value, document())


def test_generic_csv_validates_business_rules_before_resource_mutation(tmp_path):
    path = tmp_path / "readings.csv"
    path.write_text("station,amount\n0001,12.5\n", encoding="utf-8")
    _, rows = inspect_registered_source(raw_table(), tmp_path)
    assert rows == [{"station": "0001", "amount": 12.5}]
    path.write_text("station,amount\n0001,-1\n", encoding="utf-8")
    with pytest.raises(ContractBlocked, match="row 2.*minimum"):
        inspect_registered_source(raw_table(), tmp_path)
    path.write_text("station,amount\n0001,1\n0001,2\n", encoding="utf-8")
    with pytest.raises(ContractBlocked, match="duplicate"):
        inspect_registered_source(raw_table(), tmp_path)


def test_generic_excel_selects_explicit_sheet_and_preserves_formatted_identifier(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Readings"
    sheet.append(["station", "amount"])
    sheet.append([7, 12.5])
    sheet["A2"].number_format = "0000"
    workbook.create_sheet("Unrelated").append(["not", "the", "feed"])
    workbook.save(tmp_path / "readings.xlsx")
    table = raw_table(source_filename="readings.xlsx", source_sheet="Readings")
    _, rows = inspect_registered_source(table, tmp_path)
    assert rows == [{"station": "0007", "amount": 12.5}]
    table.source_sheet = None
    with pytest.raises(ContractBlocked, match="select a workbook sheet"):
        inspect_registered_source(table, tmp_path)


def test_schema_rejects_unknown_fields_invalid_metadata_and_nullable_keys():
    with pytest.raises(ValidationError):
        ContractField(
            name="x", type="STRING", nullable=False, description="test",
            evidence_quote="x is described", metadata="run_date",
        )
    with pytest.raises(ValidationError):
        raw_table(unknown="field")
    with pytest.raises(ValidationError, match="required business/source"):
        raw_table(columns=[
            ContractField(name="station", type="STRING", nullable=True,
                          description="station", evidence_quote="station STRING"),
        ])
