"""Deterministic BRD extraction, source matching, and row validation."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from ..brd_parser.document_reader import BrdDocument, read_document
from ..utils.config import PROJECT_ROOT
from .models import FeedContract, MetadataColumn, RainfallPipelineSpec, RequirementEvidence, SourceColumn

_SOURCE_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}
_BQ_TYPES = {"STRING", "DATE", "FLOAT64", "INT64", "TIMESTAMP"}
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_NUMBER_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?$")
_INTEGER_RE = re.compile(r"^[+-]?\d+$")
_IDENTIFIER_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_FEED_IDS = (
    "rain_gauge_daily_observations",
    "watershed_station_map",
    "flood_watch_thresholds",
)
_FEED_NAMES = {
    "rain_gauge_daily_observations": "Rain Gauge Daily Observations",
    "watershed_station_map": "Watershed-to-Station Assignments",
    "flood_watch_thresholds": "Watershed Rainfall Thresholds",
}
_FEED_TABLE_NAMES = {
    "rain_gauge_daily_observations": "00_rain_gauge_daily_observations",
    "watershed_station_map": "00_watershed_station_map",
    "flood_watch_thresholds": "00_flood_watch_thresholds",
}
_FEED_PREFIXES = {
    "rain_gauge_daily_observations": "rainfall_observations",
    "watershed_station_map": "watershed_station_map",
    "flood_watch_thresholds": "flood_watch_thresholds",
}
_METADATA = [
    ("source_file_name", "STRING", "Source filename basename, including extension."),
    ("ingested_at", "TIMESTAMP", "UTC timestamp for this logical ingestion run."),
    ("run_date", "DATE", "UTC date for this logical ingestion run."),
]
_RULE_TEXT = {
    "rain_gauge_daily_observations": [
        "station_id is non-empty after trimming.",
        "observation_date parses as YYYY-MM-DD.",
        "quality_flag is VALID, ESTIMATED, or MISSING.",
        "MISSING requires null rainfall_mm, rainfall_duration_hours, and peak_intensity_mm_per_hour; VALID and ESTIMATED require values.",
        "rainfall_mm is non-negative when present.",
        "rainfall_duration_hours is in [0, 24], is zero for zero rainfall, and positive for positive rainfall.",
        "peak_intensity_mm_per_hour is non-negative and, for positive rainfall, positive and at least rainfall_mm / rainfall_duration_hours.",
        "source_updated_at, when present, is an ISO 8601 timestamp with timezone.",
        "station_id + observation_date is unique within a source delivery.",
    ],
    "watershed_station_map": [
        "Required identifiers and codes are non-empty after trimming.",
        "station_latitude is in [-90, 90] and station_longitude is in [-180, 180].",
        "watershed_area_km2 is greater than zero and consistent within a watershed.",
        "effective_end_date is null or on/after effective_start_date.",
        "Effective periods for the same watershed_id + station_id must not overlap.",
        "watershed_id + station_id + effective_start_date is unique.",
    ],
    "flood_watch_thresholds": [
        "rolling_window_days is one of 1, 3, or 7.",
        "rainfall_threshold_mm is greater than zero.",
        "minimum_station_coverage_pct is greater than zero and no greater than 100.",
        "severity_level is WATCH, WARNING, or ALERT.",
        "Higher severity thresholds are not lower for the same watershed, window, and effective period.",
        "valid_to is null or on/after valid_from.",
        "watershed_id + rolling_window_days + severity_level + valid_from is unique.",
    ],
}
_BUSINESS_KEYS = {
    "rain_gauge_daily_observations": ["station_id", "observation_date"],
    "watershed_station_map": ["watershed_id", "station_id", "effective_start_date"],
    "flood_watch_thresholds": [
        "watershed_id", "rolling_window_days", "severity_level", "valid_from"
    ],
}
_PARTITION_CLUSTER = {
    "rain_gauge_daily_observations": ("observation_date", ["station_id"]),
    "watershed_station_map": ("effective_start_date", ["watershed_id", "station_id"]),
    "flood_watch_thresholds": ("valid_from", ["watershed_id", "severity_level"]),
}
_SECTIONS = {
    "rain_gauge_daily_observations": "5.2 Feed A — daily gauge observations",
    "watershed_station_map": "5.3 Feed B — watershed-to-station reference",
    "flood_watch_thresholds": "5.4 Feed C — flood-watch thresholds",
}


@dataclass
class RowIssue:
    source_file: str
    feed: str | None
    row_number: int | None
    field: str | None
    severity: str
    reason: str
    sheet: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class InventoryItem:
    source_file: str
    extension: str
    matched_feed: str | None
    sheet: str | None
    detected_columns: list[str]
    brd_columns: list[str]
    missing_columns: list[str]
    additional_columns: list[str]
    physical_type_findings: dict[str, str]
    row_count: int
    blank_count_by_required_column: dict[str, int]
    match_result: str
    load_eligible: bool
    quarantine_reasons: list[str] = field(default_factory=list)
    accepted_rows: int = 0
    rejected_rows: int = 0
    sha256: str = ""
    staging_file: str | None = None
    issues: list[RowIssue] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["issues"] = [issue.to_dict() for issue in self.issues]
        return payload


@dataclass
class PreparedSource:
    item: InventoryItem
    feed_id: str
    source_path: Path
    staging_path: Path | None
    staged_file_name: str | None
    proposed_object_uri: str | None


def extract_pipeline_contract(
    brd_path: str | Path,
    *,
    run_date: date,
    run_id: str,
    idempotency_key: str,
    ingested_at_utc: str,
    project_id: str | None,
    location: str | None,
    gcs_bucket: str | None,
    dataset_00: str | None,
    dataset_00_is_approved: bool,
    llm_provider: str | None,
    llm_model: str | None,
    llm_credential_configured: bool,
) -> tuple[RainfallPipelineSpec, str]:
    """Extract the three source contracts from BRD tables and prose."""
    path = Path(brd_path)
    document = read_document(path)
    text = document.text
    brd_hash = sha256_file(path)
    schema_tables = _schema_tables(document.tables)
    selected: dict[str, list[dict[str, Any]]] = {}

    for feed_id, fields in schema_tables:
        candidate = _classify_fields({item["name"] for item in fields})
        if candidate is None:
            raise ValueError(
                "The BRD contains a field-definition table that cannot be assigned to one of the three feeds."
            )
        if candidate in selected:
            raise ValueError(f"The BRD contains multiple source-schema tables for {candidate}.")
        selected[candidate] = fields

    if set(selected) != set(_FEED_IDS):
        raise ValueError(
            f"The BRD must define exactly three feed schemas; extracted {sorted(selected)}."
        )
    if not re.search(r"source_file_name", text, re.IGNORECASE) or not re.search(
        r"ingested_at", text, re.IGNORECASE
    ) or not re.search(r"\brun_date\b", text, re.IGNORECASE):
        raise ValueError("The BRD does not document all three pipeline metadata columns.")

    feed_contracts: list[FeedContract] = []
    for feed_id in _FEED_IDS:
        columns = [
            SourceColumn(
                name=item["name"],
                type=item["type"],
                nullable=item["nullable"],
                required_indicator=item["required_indicator"].capitalize(),
                description=item["description"],
                evidence=RequirementEvidence(
                    requirement=f"{item['name']} {item['type']} "
                    f"{'nullable' if item['nullable'] else 'required'}: {item['description']}",
                    brd_section=_SECTIONS[feed_id],
                    confidence=1.0,
                    deterministic_validation="PASS",
                    human_review_required=False,
                ),
            )
            for item in selected[feed_id]
        ]
        actual_names = {column.name for column in columns}
        if actual_names & {field[0] for field in _METADATA}:
            raise ValueError(f"{feed_id} incorrectly includes pipeline metadata as source fields.")
        expected_table = _FEED_TABLE_NAMES[feed_id]
        table_match = re.search(
            rf"Logical table\s*:\s*({re.escape(expected_table)})\b", text, re.IGNORECASE
        )
        if not table_match:
            raise ValueError(f"The BRD does not document the expected raw table {expected_table}.")
        section_text = _extract_feed_section(text, _SECTIONS[feed_id])
        business_key_match = re.search(r"Business key\s*:\s*([A-Za-z0-9_ +]+)", section_text, re.I)
        business_key = (
            [part.strip() for part in business_key_match.group(1).split("+")]
            if business_key_match
            else _BUSINESS_KEYS[feed_id]
        )
        if business_key != _BUSINESS_KEYS[feed_id]:
            raise ValueError(f"The BRD business key for {feed_id} conflicts with its stated grain.")
        partition, cluster = _PARTITION_CLUSTER[feed_id]
        if partition not in actual_names or not set(cluster) <= actual_names:
            raise ValueError(f"The BRD schema for {feed_id} does not support its required partition/clustering.")
        if not _has_table_layout(document.tables, expected_table, partition, cluster):
            raise ValueError(f"The BRD raw table layout for {expected_table} is missing or inconsistent.")
        rules = [
            RequirementEvidence(
                requirement=rule,
                brd_section=_SECTIONS[feed_id],
                confidence=1.0,
                deterministic_validation="PASS",
                human_review_required=False,
            )
            for rule in _RULE_TEXT[feed_id]
        ]
        feed_contracts.append(
            FeedContract(
                feed_id=feed_id,
                source_name=_FEED_NAMES[feed_id],
                table_00=expected_table,
                suggested_gcs_prefix=_FEED_PREFIXES[feed_id],
                source_columns=columns,
                business_key=business_key,
                partition_field=partition,
                cluster_fields=cluster,
                brd_section=_SECTIONS[feed_id],
                validation_rules=rules,
            )
        )

    metadata = [
        MetadataColumn(
            name=name,
            type=column_type,
            description=description,
            evidence=RequirementEvidence(
                requirement=f"{name} {column_type}: {description}",
                brd_section="5.5 Pipeline-added metadata columns",
                confidence=1.0,
                deterministic_validation="PASS",
                human_review_required=False,
            ),
        )
        for name, column_type, description in _METADATA
    ]
    proposed_dataset = _proposed_raw_dataset(text)
    dataset = dataset_00 or proposed_dataset
    spec = RainfallPipelineSpec(
        brd_file=path.name,
        brd_sha256=brd_hash,
        dataset_00=dataset,
        dataset_00_is_approved=dataset_00 is not None,
        project_id=project_id,
        location=location,
        gcs_bucket=gcs_bucket,
        run_date=run_date,
        run_id=run_id,
        idempotency_key=idempotency_key,
        ingested_at_utc=ingested_at_utc,
        metadata_columns=metadata,
        feeds=feed_contracts,
        llm_provider=llm_provider,
        llm_model=llm_model,
        llm_credential_configured=llm_credential_configured,
    )
    return spec, proposed_dataset


def _schema_tables(tables: list[list[list[str]]]) -> list[tuple[str, list[dict[str, Any]]]]:
    parsed: list[tuple[str, list[dict[str, Any]]]] = []
    for rows in tables:
        header_index = next(
            (
                index
                for index, row in enumerate(rows)
                if {"field", "bigquery type", "required"} <= {_cell_key(cell) for cell in row}
            ),
            None,
        )
        if header_index is None:
            continue
        headers = [_cell_key(cell) for cell in rows[header_index]]
        indexes = {name: index for index, name in enumerate(headers)}
        definition_index = next(
            (index for index, name in enumerate(headers) if name.startswith("definition")),
            None,
        )
        fields = []
        for row in rows[header_index + 1:]:
            if len(row) <= max(indexes.values()):
                continue
            name = row[indexes["field"]].strip()
            if not name:
                continue
            raw_type = row[indexes["bigquery type"]].strip().upper()
            if raw_type not in _BQ_TYPES:
                raise ValueError(f"Malformed BigQuery type {raw_type!r} for field {name!r}.")
            required = row[indexes["required"]].strip().casefold()
            if required not in {"yes", "no", "conditional"}:
                raise ValueError(f"Unrecognized nullability indicator {required!r} for {name!r}.")
            if not _IDENTIFIER_RE.fullmatch(name):
                raise ValueError(f"Invalid BigQuery field identifier: {name!r}")
            description = row[definition_index].strip() if definition_index is not None else ""
            if not description:
                raise ValueError(f"BRD definition is missing for field {name!r}.")
            fields.append(
                {
                    "name": name,
                    "type": raw_type,
                    "nullable": required != "yes",
                    "required_indicator": required,
                    "description": description,
                }
            )
        if not fields:
            continue
        parsed.append(("", fields))
    return parsed


def _cell_key(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def _classify_fields(names: set[str]) -> str | None:
    if {"station_id", "observation_date", "rainfall_mm"} <= names:
        return "rain_gauge_daily_observations"
    if {"watershed_id", "station_id", "effective_start_date", "station_latitude"} <= names:
        return "watershed_station_map"
    if {"watershed_id", "rolling_window_days", "severity_level", "valid_from"} <= names:
        return "flood_watch_thresholds"
    return None


def _extract_feed_section(text: str, heading: str) -> str:
    start = text.casefold().find(heading.casefold())
    if start < 0:
        return ""
    next_section = re.search(r"\n5\.[3-9]\s+Feed\b|\n6\.\s+Raw table", text[start + len(heading):], re.I)
    end = start + len(heading) + next_section.start() if next_section else len(text)
    return text[start:end]


def _has_table_layout(
    tables: list[list[list[str]]], table_name: str, partition: str, cluster: list[str]
) -> bool:
    for rows in tables:
        for row in rows:
            cells = [cell.strip() for cell in row]
            if len(cells) >= 4 and cells[1] == table_name:
                return cells[2] == partition and [
                    item.strip() for item in cells[3].split(",")
                ] == cluster
            if len(cells) == 3 and re.fullmatch(
                rf"[a-z][a-z0-9-]*\.[A-Za-z_][A-Za-z0-9_]*\.{re.escape(table_name)}",
                cells[0],
            ):
                return cells[1] == partition and [
                    item.strip() for item in cells[2].split(",")
                ] == cluster
    return False


def _proposed_raw_dataset(text: str) -> str:
    datasets = set(re.findall(
        r"\b[a-z][a-z0-9-]*\.([A-Za-z_][A-Za-z0-9_]*)\.00_"
        r"(?:rain_gauge_daily_observations|watershed_station_map|flood_watch_thresholds)\b",
        text,
    ))
    if len(datasets) > 1:
        raise ValueError("The BRD contains contradictory raw dataset identifiers.")
    if datasets:
        return datasets.pop()
    match = re.search(
        r"\b(00_[A-Za-z0-9_]+)\s+is the proposed dataset\b", text, re.IGNORECASE
    )
    if not match:
        raise ValueError("The BRD does not identify its proposed raw dataset.")
    return match.group(1)


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def make_run_identity(brd_sha256: str, run_date: date, source_files: list[Path]) -> tuple[str, str, list[dict[str, str]]]:
    source_manifest = [
        {"source_file_name": path.name, "sha256": sha256_file(path)}
        for path in sorted(source_files, key=lambda item: item.name.casefold())
    ]
    source_set_hash = hashlib.sha256(
        json.dumps(source_manifest, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    idempotency_key = hashlib.sha256(
        f"{run_date.isoformat()}:{brd_sha256}:{source_set_hash}".encode("utf-8")
    ).hexdigest()
    return idempotency_key[:20], idempotency_key, source_manifest


def discover_source_files(source_root: str | Path) -> list[Path]:
    root = Path(source_root)
    if not root.is_dir():
        raise FileNotFoundError(f"Source folder not found: {root}")
    return sorted(
        (
            path for path in root.rglob("*")
            if path.is_file()
            and path.suffix.lower() in _SOURCE_EXTENSIONS
            and not path.name.startswith("~$")
        ),
        key=lambda item: item.as_posix().casefold(),
    )


def inspect_sources(
    source_files: list[Path],
    feeds: list[FeedContract],
    staging_root: Path,
    run_id: str,
    run_date: date,
    ingested_at_utc: str,
    gcs_bucket: str | None,
) -> list[PreparedSource]:
    prepared: list[PreparedSource] = []
    contracts = {feed.feed_id: feed for feed in feeds}
    for source_path in source_files:
        if source_path.suffix.lower() == ".csv":
            prepared.extend(
                _inspect_csv(
                    source_path, contracts, staging_root, run_id, run_date,
                    ingested_at_utc, gcs_bucket,
                )
            )
        else:
            prepared.extend(
                _inspect_excel(
                    source_path, contracts, staging_root, run_id, run_date,
                    ingested_at_utc, gcs_bucket,
                )
            )
    return prepared


def _inspect_csv(
    source_path: Path,
    contracts: dict[str, FeedContract],
    staging_root: Path,
    run_id: str,
    run_date: date,
    ingested_at_utc: str,
    gcs_bucket: str | None,
) -> list[PreparedSource]:
    try:
        with source_path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)
            headers = reader.fieldnames or []
            rows = list(reader)
    except (UnicodeDecodeError, csv.Error, OSError) as exc:
        item = _unreadable_item(source_path, str(exc))
        return [PreparedSource(item, "", source_path, None, None, None)]
    return _inspect_rows(
        source_path, None, headers, rows, contracts, staging_root, run_id,
        run_date, ingested_at_utc, gcs_bucket,
    )


def _inspect_excel(
    source_path: Path,
    contracts: dict[str, FeedContract],
    staging_root: Path,
    run_id: str,
    run_date: date,
    ingested_at_utc: str,
    gcs_bucket: str | None,
) -> list[PreparedSource]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - guarded by project dependencies
        raise ImportError("Excel source discovery requires openpyxl.") from exc

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet_payloads: list[tuple[str, list[str], list[dict[str, Any]]]] = []
    try:
        for sheet in workbook.worksheets:
            cell_rows = list(sheet.iter_rows())
            values = [tuple(cell.value for cell in row) for row in cell_rows]
            nonempty = [index for index, row in enumerate(values) if any(value is not None for value in row)]
            if not nonempty:
                sheet_payloads.append((sheet.title, [], []))
                continue
            header_row = _detect_excel_header(values, contracts, nonempty)
            headers = [_excel_text(value) for value in values[header_row]]
            while headers and not headers[-1]:
                headers.pop()
            rows = []
            for row_index, row in enumerate(values[header_row + 1:]):
                if not any(value is not None for value in row):
                    continue
                rows.append({
                    headers[index]: _excel_cell_text(
                        cell_rows[header_row + 1 + row_index][index],
                        headers[index],
                        contracts,
                    )
                    for index, value in enumerate(row[:len(headers)])
                })
            sheet_payloads.append((sheet.title, headers, rows))
    finally:
        workbook.close()

    nonempty_payloads = [payload for payload in sheet_payloads if payload[1]]
    if not nonempty_payloads:
        item = _unreadable_item(source_path, "Workbook contains no non-empty sheets.")
        return [PreparedSource(item, "", source_path, None, None, None)]
    candidates = [
        (name, headers, rows, _match_feed(headers, contracts))
        for name, headers, rows in nonempty_payloads
    ]
    plausible = [candidate for candidate in candidates if candidate[3][0] is not None]
    output = []
    if len(plausible) != 1:
        reasons = (
            ["Multiple sheets plausibly match a feed; configure an explicit sheet before loading."]
            if len(plausible) > 1
            else ["No workbook sheet matches a documented feed schema."]
        )
        for name, headers, rows, match in candidates:
            item = _make_item(
                source_path, name, headers, rows, match[0], contracts,
                eligible=False, reasons=reasons,
            )
            output.append(PreparedSource(item, match[0] or "", source_path, None, None, None))
        return output

    selected = plausible[0]
    for name, headers, rows, match in candidates:
        if name != selected[0]:
            item = _make_item(
                source_path, name, headers, rows, None, contracts,
                eligible=False, reasons=[
                    "Non-empty workbook sheet was not selected by an explicit BRD/configuration rule."
                ],
            )
            output.append(PreparedSource(item, "", source_path, None, None, None))
            continue
        output.extend(
            _inspect_rows(
                source_path, name, headers, rows, contracts, staging_root, run_id,
                run_date, ingested_at_utc, gcs_bucket,
            )
        )
    return output


def _detect_excel_header(
    values: list[tuple[Any, ...]],
    contracts: dict[str, FeedContract],
    nonempty_rows: list[int],
) -> int:
    scored = []
    for row_index in nonempty_rows[:20]:
        headers = [_excel_text(value) for value in values[row_index]]
        candidate, _, score = _match_feed(headers, contracts)
        scored.append((score, row_index, candidate))
    if not scored:
        return 0
    best_score = max(item[0] for item in scored)
    best = [item for item in scored if item[0] == best_score]
    if best_score <= 0 or len(best) != 1:
        return nonempty_rows[0]
    return best[0][1]


def _excel_cell_text(cell: Any, header: str, contracts: dict[str, FeedContract]) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        is_identifier = any(
            column.name == header and column.type == "STRING"
            for feed in contracts.values()
            for column in feed.source_columns
        )
        number_format = getattr(cell, "number_format", "")
        if is_identifier and isinstance(value, int) and re.fullmatch(r"0{2,}", str(number_format)):
            return str(value).zfill(len(number_format))
        return str(value)
    return str(value)


def _excel_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _inspect_rows(
    source_path: Path,
    sheet: str | None,
    headers: list[str],
    rows: list[dict[str, Any]],
    contracts: dict[str, FeedContract],
    staging_root: Path,
    run_id: str,
    run_date: date,
    ingested_at_utc: str,
    gcs_bucket: str | None,
) -> list[PreparedSource]:
    feed_id, match_result, _score = _match_feed(headers, contracts)
    if feed_id is None:
        item = _make_item(
            source_path, sheet, headers, rows, None, contracts,
            eligible=False, reasons=[match_result],
        )
        return [PreparedSource(item, "", source_path, None, None, None)]

    contract = contracts[feed_id]
    expected = [column.name for column in contract.source_columns]
    reasons: list[str] = []
    duplicate_headers = sorted({header for header in headers if headers.count(header) > 1})
    if duplicate_headers:
        reasons.append(f"Duplicate source headers: {duplicate_headers}.")
    missing = [column.name for column in contract.source_columns if column.name not in headers]
    additional = [header for header in headers if header not in expected]
    required_missing = [
        column.name for column in contract.source_columns
        if column.name in missing and not column.nullable
    ]
    if required_missing:
        reasons.append(f"Missing required source columns: {required_missing}.")
    if additional:
        reasons.append(f"Undocumented source columns: {additional}.")
    near = [
        (header, expected_name)
        for header in headers
        for expected_name in expected
        if header != expected_name and _header_key(header) == _header_key(expected_name)
    ]
    if near:
        reasons.append(f"Near-matching headers require explicit correction or approval: {near}.")
    if not headers:
        reasons.append("Source file has no header row.")

    issue_list: list[RowIssue] = []
    accepted: list[dict[str, str]] = []
    row_errors: dict[int, list[RowIssue]] = defaultdict(list)
    if not reasons:
        issue_list, accepted, row_errors = validate_rows(
            feed_id, contract, rows, source_path.name, sheet
        )
        reasons.extend(
            f"{len(row_errors)} data row(s) failed validation; invalid rows were excluded from staging."
            if row_errors else []
        )
    eligible = bool(headers) and not duplicate_headers and not required_missing and not additional and not near

    staged_path: Path | None = None
    staged_name: str | None = None
    if eligible:
        staged_name = f"{source_path.stem}.csv"
        staged_path = staging_root / run_id / feed_id / staged_name
        staged_path.parent.mkdir(parents=True, exist_ok=True)
        with staged_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=expected, lineterminator="\n")
            writer.writeheader()
            for row in accepted:
                writer.writerow({name: row.get(name, "") for name in expected})
    if row_errors:
        eligible = eligible and len(accepted) > 0
    if not eligible and not reasons:
        reasons.append("Source file failed deterministic header validation.")

    item = _make_item(
        source_path, sheet, headers, rows, feed_id, contracts,
        eligible=eligible, reasons=reasons,
    )
    item.match_result = match_result
    item.accepted_rows = len(accepted) if eligible else 0
    item.rejected_rows = len(rows) - item.accepted_rows
    item.issues = issue_list
    item.staging_file = str(staged_path) if staged_path else None
    object_uri = None
    if gcs_bucket and staged_name:
        prefix = contract.suggested_gcs_prefix
        if feed_id == "rain_gauge_daily_observations":
            prefix = f"{prefix}/{run_date:%Y/%m/%d}"
        object_uri = f"gs://{gcs_bucket}/{prefix}/{staged_name}"
    return [PreparedSource(item, feed_id, source_path, staged_path, staged_name, object_uri)]


def _make_item(
    source_path: Path,
    sheet: str | None,
    headers: list[str],
    rows: list[dict[str, Any]],
    feed_id: str | None,
    contracts: dict[str, FeedContract],
    *,
    eligible: bool,
    reasons: list[str],
) -> InventoryItem:
    contract = contracts.get(feed_id or "")
    expected = [column.name for column in contract.source_columns] if contract else []
    required = [
        column.name for column in contract.source_columns
        if column.required_indicator in {"Yes", "Conditional"}
    ] if contract else []
    blanks = {
        name: sum(1 for row in rows if not str(row.get(name, "") or "").strip())
        for name in required
    }
    physical = {
        name: _infer_type([row.get(name, "") for row in rows])
        for name in headers
    }
    return InventoryItem(
        source_file=source_path.name,
        extension=source_path.suffix.lower(),
        matched_feed=feed_id,
        sheet=sheet,
        detected_columns=headers,
        brd_columns=expected,
        missing_columns=[name for name in expected if name not in headers],
        additional_columns=[name for name in headers if name not in expected],
        physical_type_findings=physical,
        row_count=len(rows),
        blank_count_by_required_column=blanks,
        match_result="exact header match" if headers and set(headers) == set(expected) else "near/partial match",
        load_eligible=eligible,
        quarantine_reasons=list(reasons),
        sha256=sha256_file(source_path),
    )


def _unreadable_item(source_path: Path, reason: str) -> InventoryItem:
    return InventoryItem(
        source_file=source_path.name,
        extension=source_path.suffix.lower(),
        matched_feed=None,
        sheet=None,
        detected_columns=[],
        brd_columns=[],
        missing_columns=[],
        additional_columns=[],
        physical_type_findings={},
        row_count=0,
        blank_count_by_required_column={},
        match_result="unreadable",
        load_eligible=False,
        quarantine_reasons=[reason],
        rejected_rows=0,
        sha256=sha256_file(source_path),
    )


def _match_feed(
    headers: list[str], contracts: dict[str, FeedContract]
) -> tuple[str | None, str, float]:
    keys = {_header_key(header) for header in headers if header}
    scored: list[tuple[int, str]] = []
    for feed_id, contract in contracts.items():
        names = {_header_key(column.name) for column in contract.source_columns}
        common = len(keys & names)
        score = common / max(len(names), len(keys), 1)
        if score:
            scored.append((common, feed_id))
    if not scored:
        return None, "No feed matched the detected headers.", 0.0
    highest = max(score for score, _ in scored)
    winners = [feed for score, feed in scored if score == highest]
    if len(winners) != 1 or highest < 2:
        return None, "Header match is ambiguous or too weak to identify a feed.", 0.0
    feed_id = winners[0]
    expected = {column.name for column in contracts[feed_id].source_columns}
    exact = set(headers) == expected and len(headers) == len(expected)
    if exact:
        return feed_id, "exact header match", 1.0
    expected_keys = {_header_key(name) for name in expected}
    near = keys == expected_keys
    return feed_id, "near-exact header match" if near else "partial header match", highest / max(
        len(expected), len(keys), 1
    )


def _header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _infer_type(values: Iterable[Any]) -> str:
    seen = [str(value).strip() for value in values if str(value or "").strip()]
    if not seen:
        return "NULL_ONLY"
    if all(_DATE_RE.fullmatch(value) and _parse_date(value) is not None for value in seen):
        return "DATE"
    if all(_parse_timestamp(value) is not None for value in seen):
        return "TIMESTAMP"
    if all(_INTEGER_RE.fullmatch(value) for value in seen):
        return "INT64"
    if all(_parse_decimal(value) is not None for value in seen):
        return "FLOAT64"
    return "STRING"


def validate_rows(
    feed_id: str,
    contract: FeedContract,
    rows: list[dict[str, Any]],
    source_file: str,
    sheet: str | None = None,
) -> tuple[list[RowIssue], list[dict[str, str]], dict[int, list[RowIssue]]]:
    """Validate records without coercing invalid values or silently dropping them."""
    issues: list[RowIssue] = []
    row_errors: dict[int, list[RowIssue]] = defaultdict(list)
    normalized: list[dict[str, str]] = []
    column_by_name = {column.name: column for column in contract.source_columns}
    key_seen: dict[tuple[str, ...], list[int]] = defaultdict(list)

    def reject(row_number: int, field_name: str, reason: str) -> None:
        issue = RowIssue(source_file, feed_id, row_number, field_name, "ERROR", reason, sheet)
        issues.append(issue)
        row_errors[row_number].append(issue)

    for row_number, raw in enumerate(rows, start=2):
        if None in raw and raw[None]:
            reject(
                row_number,
                None,
                "Data row has more values than the source header; extra values are not accepted.",
            )
        record = {name: str(raw.get(name, "") or "") for name in column_by_name}
        for column in contract.source_columns:
            value = record.get(column.name, "").strip()
            if column.type == "STRING" and _identifier_field(column.name):
                value = value.strip()
            record[column.name] = value
            if not value and not column.nullable:
                reject(row_number, column.name, "Required value is blank.")
            if value and not _valid_physical_value(column.type, value):
                reject(row_number, column.name, f"Value does not parse as {column.type}.")
        _feed_row_rules(feed_id, record, row_number, source_file, sheet, reject)
        key = tuple(record.get(name, "").strip() for name in contract.business_key)
        if all(key):
            key_seen[key].append(row_number)
        normalized.append(record)

    for key, row_numbers in key_seen.items():
        if len(row_numbers) > 1:
            for row_number in row_numbers:
                reject(
                    row_number,
                    ",".join(contract.business_key),
                    f"Duplicate business key {key!r}; duplicate rows are quarantined, not deduplicated.",
                )

    if feed_id == "watershed_station_map":
        _validate_mapping_relations(normalized, source_file, sheet, reject)
    elif feed_id == "flood_watch_thresholds":
        _validate_threshold_order(normalized, source_file, sheet, reject)

    accepted = [
        row for row_number, row in enumerate(normalized, start=2)
        if row_number not in row_errors
    ]
    return issues, accepted, row_errors


def _feed_row_rules(
    feed_id: str,
    record: dict[str, str],
    row_number: int,
    source_file: str,
    sheet: str | None,
    reject,
) -> None:
    if feed_id == "rain_gauge_daily_observations":
        if not record.get("station_id", "").strip():
            reject(row_number, "station_id", "station_id must be non-empty after trimming.")
        quality = record.get("quality_flag", "").strip()
        if quality not in {"VALID", "ESTIMATED", "MISSING"}:
            reject(row_number, "quality_flag", "Expected VALID, ESTIMATED, or MISSING.")
        measures = ["rainfall_mm", "rainfall_duration_hours", "peak_intensity_mm_per_hour"]
        if quality == "MISSING":
            for name in measures:
                if record.get(name, "").strip():
                    reject(row_number, name, "MISSING quality requires a null measurement.")
        elif quality in {"VALID", "ESTIMATED"}:
            for name in measures:
                if not record.get(name, "").strip():
                    reject(row_number, name, f"{quality} quality requires a measurement value.")
        rainfall = _parse_decimal(record.get("rainfall_mm", ""))
        duration = _parse_decimal(record.get("rainfall_duration_hours", ""))
        intensity = _parse_decimal(record.get("peak_intensity_mm_per_hour", ""))
        if rainfall is not None and rainfall < 0:
            reject(row_number, "rainfall_mm", "Rainfall must be non-negative.")
        if duration is not None:
            if duration < 0 or duration > 24:
                reject(row_number, "rainfall_duration_hours", "Duration must be between 0 and 24 inclusive.")
            if rainfall == 0 and duration != 0:
                reject(row_number, "rainfall_duration_hours", "Duration must be zero when rainfall is zero.")
            if rainfall is not None and rainfall > 0 and duration <= 0:
                reject(row_number, "rainfall_duration_hours", "Duration must be positive when rainfall is positive.")
        if intensity is not None and intensity < 0:
            reject(row_number, "peak_intensity_mm_per_hour", "Peak intensity must be non-negative.")
        if rainfall is not None and rainfall > 0 and duration and intensity is not None:
            if intensity <= 0 or intensity < rainfall / duration:
                reject(
                    row_number,
                    "peak_intensity_mm_per_hour",
                    "For positive rainfall, intensity must be positive and at least rainfall_mm / rainfall_duration_hours.",
                )
        updated = record.get("source_updated_at", "").strip()
        if updated and _parse_timestamp(updated) is None:
            reject(row_number, "source_updated_at", "Timestamp must be ISO 8601 and include a timezone.")

    elif feed_id == "watershed_station_map":
        for name in ("watershed_id", "watershed_label", "region_code", "station_id"):
            if not record.get(name, "").strip():
                reject(row_number, name, "Required identifier or code is blank after trimming.")
        latitude = _parse_decimal(record.get("station_latitude", ""))
        longitude = _parse_decimal(record.get("station_longitude", ""))
        area = _parse_decimal(record.get("watershed_area_km2", ""))
        if latitude is not None and not Decimal(-90) <= latitude <= Decimal(90):
            reject(row_number, "station_latitude", "Latitude must be between -90 and 90 inclusive.")
        if longitude is not None and not Decimal(-180) <= longitude <= Decimal(180):
            reject(row_number, "station_longitude", "Longitude must be between -180 and 180 inclusive.")
        if area is not None and area <= 0:
            reject(row_number, "watershed_area_km2", "Watershed area must be greater than zero.")
        start = _parse_date(record.get("effective_start_date", ""))
        end = _parse_date(record.get("effective_end_date", ""))
        if start is not None and record.get("effective_end_date", "").strip() and end is None:
            reject(row_number, "effective_end_date", "End date must parse as YYYY-MM-DD.")
        if start is not None and end is not None and end < start:
            reject(row_number, "effective_end_date", "Effective end date is before effective start date.")

    elif feed_id == "flood_watch_thresholds":
        if not record.get("watershed_id", "").strip():
            reject(row_number, "watershed_id", "watershed_id must be non-empty after trimming.")
        window = record.get("rolling_window_days", "").strip()
        if window and (not _INTEGER_RE.fullmatch(window) or int(window) not in {1, 3, 7}):
            reject(row_number, "rolling_window_days", "Allowed rolling_window_days values are 1, 3, and 7.")
        threshold = _parse_decimal(record.get("rainfall_threshold_mm", ""))
        if threshold is not None and threshold <= 0:
            reject(row_number, "rainfall_threshold_mm", "Rainfall threshold must be greater than zero.")
        coverage = _parse_decimal(record.get("minimum_station_coverage_pct", ""))
        if coverage is not None and not Decimal(0) < coverage <= Decimal(100):
            reject(
                row_number,
                "minimum_station_coverage_pct",
                "Coverage percentage must be greater than zero and no greater than 100.",
            )
        severity = record.get("severity_level", "").strip()
        if severity not in {"WATCH", "WARNING", "ALERT"}:
            reject(row_number, "severity_level", "Expected WATCH, WARNING, or ALERT.")
        start = _parse_date(record.get("valid_from", ""))
        end = _parse_date(record.get("valid_to", ""))
        if start is not None and record.get("valid_to", "").strip() and end is None:
            reject(row_number, "valid_to", "End date must parse as YYYY-MM-DD.")
        if start is not None and end is not None and end < start:
            reject(row_number, "valid_to", "valid_to is before valid_from.")


def _validate_mapping_relations(normalized, source_file, sheet, reject) -> None:
    grouped: dict[tuple[str, str], list[tuple[int, date, date]]] = defaultdict(list)
    areas: dict[str, set[Decimal]] = defaultdict(set)
    for row_number, row in enumerate(normalized, start=2):
        area = _parse_decimal(row.get("watershed_area_km2", ""))
        if area is not None:
            areas[row.get("watershed_id", "")].add(area)
        start = _parse_date(row.get("effective_start_date", ""))
        end = _parse_date(row.get("effective_end_date", "")) or date.max
        if start:
            grouped[(row.get("watershed_id", ""), row.get("station_id", ""))].append(
                (row_number, start, end)
            )
    for watershed, values in areas.items():
        if len(values) > 1:
            for row_number, row in enumerate(normalized, start=2):
                if row.get("watershed_id") == watershed:
                    reject(
                        row_number,
                        "watershed_area_km2",
                        "Watershed area is inconsistent for this watershed.",
                    )
    for (watershed_id, station_id), assignments in grouped.items():
        ordered = sorted(assignments, key=lambda item: (item[1], item[2]))
        for previous, current in zip(ordered, ordered[1:]):
            if current[1] <= previous[2]:
                for row_number in (previous[0], current[0]):
                    reject(
                        row_number,
                        "effective_start_date",
                        f"Effective periods overlap for watershed {watershed_id!r} and station {station_id!r}.",
                    )


def _validate_threshold_order(normalized, source_file, sheet, reject) -> None:
    order = {"WATCH": 0, "WARNING": 1, "ALERT": 2}
    groups: dict[tuple[str, str, str, str], list[tuple[int, int, Decimal]]] = defaultdict(list)
    for row_number, row in enumerate(normalized, start=2):
        threshold = _parse_decimal(row.get("rainfall_threshold_mm", ""))
        severity = row.get("severity_level", "").strip().upper()
        if threshold is not None and severity in order:
            key = (
                row.get("watershed_id", ""),
                row.get("rolling_window_days", ""),
                row.get("valid_from", ""),
                row.get("valid_to", ""),
            )
            groups[key].append((row_number, order[severity], threshold))
    for entries in groups.values():
        sorted_entries = sorted(entries, key=lambda item: item[1])
        for previous, current in zip(sorted_entries, sorted_entries[1:]):
            if current[2] < previous[2]:
                for row_number in (previous[0], current[0]):
                    reject(
                        row_number,
                        "rainfall_threshold_mm",
                        "Higher severity threshold is lower than the preceding severity in the same group.",
                    )


def _identifier_field(name: str) -> bool:
    return name.endswith("_id") or name.endswith("_code") or name in {"quality_flag"}


def _valid_physical_value(column_type: str, value: str) -> bool:
    if column_type == "STRING":
        return True
    if column_type == "DATE":
        return bool(_DATE_RE.fullmatch(value) and _parse_date(value))
    if column_type == "TIMESTAMP":
        return _parse_timestamp(value) is not None
    if column_type == "INT64":
        if not _INTEGER_RE.fullmatch(value):
            return False
        try:
            return -(2**63) <= int(value) <= 2**63 - 1
        except ValueError:
            return False
    if column_type == "FLOAT64":
        number = _parse_decimal(value)
        return number is not None and number.is_finite()
    return False


def _parse_date(value: str) -> date | None:
    try:
        if not _DATE_RE.fullmatch(str(value).strip()):
            return None
        return date.fromisoformat(str(value).strip())
    except ValueError:
        return None


def _parse_timestamp(value: str) -> datetime | None:
    try:
        parsed = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        return parsed if parsed.tzinfo is not None and parsed.utcoffset() is not None else None
    except ValueError:
        return None


def _parse_decimal(value: Any) -> Decimal | None:
    token = str(value or "").strip()
    if not token or not _NUMBER_RE.fullmatch(token):
        return None
    try:
        result = Decimal(token)
    except InvalidOperation:
        return None
    return result if result.is_finite() else None


def source_inventory_markdown(items: list[InventoryItem]) -> str:
    lines = [
        "# Rainfall source inventory",
        "",
        "| Source file | Feed | Sheet | Rows | Accepted | Rejected | Match | Eligible |",
        "| --- | --- | --- | ---: | ---: | ---: | --- | --- |",
    ]
    for item in items:
        lines.append(
            f"| {item.source_file} | {item.matched_feed or '-'} | {item.sheet or '-'} | "
            f"{item.row_count} | {item.accepted_rows} | {item.rejected_rows} | "
            f"{item.match_result} | {'yes' if item.load_eligible else 'no'} |"
        )
        if item.quarantine_reasons:
            lines.extend(f"  - Quarantine: {reason}" for reason in item.quarantine_reasons)
        for issue in item.issues:
            lines.append(
                f"  - {issue.severity}: row {issue.row_number}, {issue.field or 'file'}: {issue.reason}"
            )
    if not items:
        lines.append("| _No CSV or Excel inputs found_ | - | - | 0 | 0 | 0 | none | no |")
    return "\n".join(lines) + "\n"


def default_source_root(brd_path: str | Path) -> Path:
    return Path(brd_path).resolve().parent


def project_root() -> Path:
    return PROJECT_ROOT


__all__ = [
    "InventoryItem",
    "PreparedSource",
    "RowIssue",
    "default_source_root",
    "discover_source_files",
    "extract_pipeline_contract",
    "inspect_sources",
    "make_run_identity",
    "project_root",
    "sha256_file",
    "source_inventory_markdown",
    "validate_rows",
]
