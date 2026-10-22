"""BRD-driven managed-table CLI: supported rainfall compiler or generic validated SELECT plan."""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from google.api_core.exceptions import GoogleAPICallError, NotFound
from google.auth.exceptions import GoogleAuthError
from google.cloud import bigquery, storage
from sqlglot import exp, parse_one

from ..brd_parser.document_reader import read_document
from ..llm.provider import LLMError, get_provider
from ..models.pipeline_spec import ColumnSpec
from ..rainfall.executor import (
    ContractBlocked, LABELS, ensure_object, execute as execute_rainfall, field_schema,
    load_rows, schema_actions, typed_record, verify_unchanged_retry, write_report,
)
from ..rainfall.preflight import authentication_guidance, verify_workload_identity
from ..rainfall.processing import make_run_identity, sha256_file
from ..sql_generator.render_template import SqlRenderer, write_sql
from ..sql_generator.dsp import compile_dsp_publication
from ..utils.config import PROJECT_ROOT
from .automatic_runner import validate_value
from .managed_contract import ManagedContract, ManagedTable, propose_contract


def inspect_registered_source(table: ManagedTable, root: Path) -> tuple[Path, list[dict]]:
    candidates = list(root.rglob(table.source_filename))
    if len(candidates) != 1:
        raise ContractBlocked(f"{table.source_filename}: missing or ambiguous registered source.")
    path = candidates[0]
    fields = [column for column in table.columns if column.metadata is None]
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            records = list(csv.reader(handle, strict=True))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        if not table.source_sheet:
            raise ContractBlocked(f"{path.name}: BRD must explicitly select a workbook sheet.")
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if table.source_sheet not in workbook.sheetnames:
                raise ContractBlocked(f"{path.name}: documented sheet is absent.")
            rows = list(workbook[table.source_sheet].iter_rows())
            records = []
            for index, cells in enumerate(rows):
                values = []
                for position, cell in enumerate(cells):
                    value = cell.value
                    field = fields[position] if index and position < len(fields) else None
                    if isinstance(value, datetime):
                        if field and field.type == "DATE":
                            if value.time().isoformat() != "00:00:00":
                                raise ContractBlocked("A DATE cell contains a non-midnight time.")
                            value = value.date().isoformat()
                        else:
                            value = value.isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    elif field and field.type == "STRING" and isinstance(value, int):
                        if cell.number_format and set(cell.number_format) == {"0"}:
                            value = str(value).zfill(len(cell.number_format))
                    values.append("" if value is None else str(value))
                if any(values):
                    records.append(values)
        finally:
            workbook.close()
    else:
        raise ContractBlocked("Unsupported registered source format; document CSV/XLSX conversion.")
    if not records or records[0] != [column.name for column in fields]:
        raise ContractBlocked(f"{path.name}: exact ordered source headers differ from BRD.")
    accepted, keys = [], set()
    for row_number, values in enumerate(records[1:], start=2):
        if len(values) != len(fields):
            raise ContractBlocked(f"{path.name}: row {row_number} has a mismatched field count.")
        row = dict(zip([column.name for column in fields], values))
        for column in fields:
            try:
                validate_value(row[column.name], ColumnSpec(
                    name=column.name, type=column.type, nullable=column.nullable,
                ))
            except ValueError:
                raise ContractBlocked(
                    f"{path.name}: row {row_number}, field {column.name}: invalid {column.type}/nullability."
                ) from None
        for rule in table.row_rules:
            value = row.get(rule.field, "")
            if rule.operator == "nonempty":
                valid = bool(value.strip())
            elif not value:
                continue
            elif rule.operator == "allowed_values":
                valid = value in rule.values
            elif rule.operator == "date_not_before":
                valid = not row.get(rule.other_field) or date.fromisoformat(value) >= date.fromisoformat(row[rule.other_field])
            else:
                number, bound = Decimal(value), Decimal(rule.values[0])
                valid = {
                    "minimum": number >= bound, "maximum": number <= bound, "greater_than": number > bound,
                }[rule.operator]
            if not valid:
                raise ContractBlocked(f"{path.name}: row {row_number}, field {rule.field}: {rule.operator} failed.")
        key = tuple(row[name] for name in table.business_keys)
        if key in keys:
            raise ContractBlocked(f"{path.name}: duplicate source business key at row {row_number}.")
        keys.add(key)
        accepted.append(typed_record(row, fields))
    if not accepted:
        raise ContractBlocked(f"{path.name}: empty input cannot satisfy this batch.")
    return path, accepted


def compile_publication(contract: ManagedContract) -> str:
    builds, merges = [], []
    temp_names = {table.table_id: f"build_{index}" for index, table in enumerate(contract.tables) if table.layer == "01"}
    for table in contract.tables:
        if table.layer != "01":
            continue
        tree = parse_one(table.select_sql, read="bigquery")
        for reference in tree.find_all(exp.Table):
            identifier = ".".join((reference.catalog, reference.db, reference.name))
            if identifier in temp_names:
                reference.set("this", exp.to_identifier(temp_names[identifier]))
                reference.set("db", None)
                reference.set("catalog", None)
        source = temp_names[table.table_id]
        builds.append(f"CREATE TEMP TABLE {source} AS {tree.sql(dialect='bigquery')};")
        keys = ", ".join(f"`{key}`" for key in table.business_keys)
        builds.append(
            f"ASSERT NOT EXISTS (SELECT {keys} FROM {source} GROUP BY {keys} HAVING COUNT(*) > 1) "
            "AS 'FAILED: duplicate output keys';"
        )
        required = [f"`{column.name}` IS NULL" for column in table.columns if not column.nullable]
        if required:
            builds.append(
                f"ASSERT NOT EXISTS (SELECT 1 FROM {source} WHERE {' OR '.join(required)}) "
                "AS 'FAILED: required output fields';"
            )
        names = [column.name for column in table.columns]
        join = " AND ".join(f"t.`{key}` = s.`{key}`" for key in table.business_keys)
        comparable = [column.name for column in table.columns if not column.metadata]
        t = ", ".join(f"t.`{name}`" for name in comparable)
        s = ", ".join(f"s.`{name}`" for name in comparable)
        update = ", ".join(f"`{name}` = s.`{name}`" for name in names if name not in table.business_keys)
        matched = (
            f"WHEN MATCHED AND TO_JSON_STRING(STRUCT({t})) != TO_JSON_STRING(STRUCT({s})) "
            f"THEN UPDATE SET {update}\n" if update and comparable else ""
        )
        merges.append(
            f"MERGE `{table.table_id}` t USING {source} s ON {join}\n{matched}"
            f"WHEN NOT MATCHED THEN INSERT ({', '.join('`' + name + '`' for name in names)}) "
            f"VALUES ({', '.join('s.`' + name + '`' for name in names)});"
        )
    return "\n\n".join([*builds, "BEGIN TRANSACTION;", *merges, "COMMIT TRANSACTION;"])


def execute_generic(
    brd: Path, source_root: Path, output_root: Path, *, dry_run: bool, verify_retry: bool = False,
) -> dict:
    now = datetime.now(timezone.utc)
    directory = output_root / "reports" / f"generic_{now:%Y%m%dT%H%M%S%f}"
    report = {
        "status": "RUNNING", "brd_file": brd.name, "steps": [],
        "issues": [], "row_counts": {}, "jobs": [], "objects": [], "cloud_mutations": [],
    }
    phase = "authentication"
    try:
        auth = verify_workload_identity()
        report["authentication"] = auth
        if auth["status"] != "PASS":
            raise ContractBlocked(authentication_guidance(auth))
        report["steps"].append({"step": phase, "actor": "Python/Google ADC", "status": "PASS"})
        report["brd_sha256"] = sha256_file(brd)
        phase = "LLM BRD interpretation"
        contract, response = propose_contract(
            read_document(brd), get_provider(os.getenv("LLM_PROVIDER"), model=os.getenv("LLM_MODEL", "")),
            expected_project=os.getenv("GOOGLE_CLOUD_PROJECT"),
        )
        report["llm"] = {"response": response, "contract": contract.model_dump()}
        if contract.unresolved or contract.unsupported:
            for item in [*contract.unresolved, *contract.unsupported]:
                report["issues"].append({"check": "LLM contract checklist", "result": item,
                                         "action": "Clarify the BRD or implement the missing capability, then rerun."})
            raise ContractBlocked("BRD has unresolved/unsupported requirements; see the recorded LLM contract.")
        report["steps"].append({"step": phase, "actor": "LLM + Python schema/evidence/SQL AST validator",
                                "status": "PASS", "detail": contract.execution_summary})
        phase = "source validation"
        prepared = {
            table.table_id: inspect_registered_source(table, source_root)
            for table in contract.tables if table.layer == "00"
        }
        report["steps"].append({"step": phase, "actor": "Python types/keys/row rules", "status": "PASS",
                                "detail": f"{sum(len(rows) for _, rows in prepared.values())} accepted records."})
        run_id, key, manifest = make_run_identity(report["brd_sha256"], now.date(), [
            source[0] for source in prepared.values()
        ])
        directory = output_root / "reports" / f"generic_{run_id}"
        report.update(run_id=run_id, idempotency_key=key, source_manifest=manifest)
        bq = bigquery.Client(project=contract.project, location=contract.location)
        gcs = storage.Client(project=contract.project)
        phase = "cloud schema inspection"
        wanted = []
        for table in contract.tables:
            target = bigquery.Table(table.table_id, schema=[field_schema(column) for column in table.columns])
            if table.partition_by:
                target.time_partitioning = bigquery.TimePartitioning(
                    field=table.partition_by, type_=table.partition_granularity,
                )
            target.clustering_fields = table.cluster_by or None
            target.labels = LABELS
            try:
                existing = bq.get_table(table.table_id)
            except NotFound:
                additions = None
            else:
                additions = schema_actions(existing, target, allow_nullable_additions=table.layer == "00")
            wanted.append((table, target, additions))
            sql = SqlRenderer().render("managed_table.sql.j2", {
                "table_id": table.table_id, "brd_sha256": report["brd_sha256"],
                "columns": [column.model_dump() for column in table.columns],
                "partition_by": table.partition_by, "cluster_by": table.cluster_by,
                "partition_granularity": table.partition_granularity,
            })
            write_sql(sql, output_root / "sql" / table.layer, target.table_id + ".sql")
        datasets = {table.table_id.rsplit(".", 1)[0] for table in contract.tables}
        for identifier in datasets:
            try:
                dataset = bq.get_dataset(identifier)
            except NotFound:
                continue
            if dataset.location.upper() != contract.location:
                raise ContractBlocked("Dataset location conflicts with BRD.")
        try:
            bucket = gcs.get_bucket(contract.bucket)
        except NotFound:
            bucket = None
        if bucket and (
            bucket.location.upper() != contract.location
            or not bucket.iam_configuration.uniform_bucket_level_access_enabled
            or bucket.iam_configuration.public_access_prevention != "enforced"
        ):
            raise ContractBlocked("Bucket does not meet required location/private-access protections.")
        transform = compile_publication(contract)
        write_sql(transform, output_root / "sql" / "01", f"managed_transform_{run_id}.sql")
        kpi_tables = [table for table in contract.tables if table.layer == "02"]
        kpi_transform = compile_dsp_publication(kpi_tables) if kpi_tables else None
        if kpi_transform:
            write_sql(kpi_transform, output_root / "sql" / "02", f"managed_kpi_{run_id}.sql")
        report["steps"].append({"step": phase, "actor": "Python/GCP read-only inspection", "status": "PASS"})
        if dry_run:
            report["status"] = "DRY_RUN_READY"
            return report
        phase = "resource provisioning"
        if bucket is None:
            bucket = storage.Bucket(gcs, contract.bucket)
            bucket.iam_configuration.uniform_bucket_level_access_enabled = True
            bucket.iam_configuration.public_access_prevention = "enforced"
            bucket.labels = LABELS
            bucket = gcs.create_bucket(bucket, location=contract.location)
        report["objects"].append(ensure_object(
            bucket, f"brd/{report['brd_sha256']}/{brd.name}", brd.read_bytes(), content_type="application/octet-stream",
        ))
        for identifier in datasets:
            dataset = bigquery.Dataset(identifier)
            dataset.location, dataset.labels = contract.location, LABELS
            bq.create_dataset(dataset, exists_ok=True)
        for table, target, additions in wanted:
            if additions is None:
                bq.create_table(target)
            elif additions:
                for field in additions:
                    bq.query(f"ALTER TABLE `{table.table_id}` ADD COLUMN IF NOT EXISTS `{field.name}` {field.field_type}").result()
        report["steps"].append({"step": phase, "actor": "Python/GCS/BigQuery", "status": "PASS",
                                "detail": f"{len(wanted)} native managed tables created or checked."})
        params = [
            bigquery.ScalarQueryParameter("run_date", "DATE", now.date()),
            bigquery.ScalarQueryParameter("transformed_at", "TIMESTAMP", now),
        ]
        phase = "BigQuery SELECT schema validation"
        for table in contract.tables:
            if table.layer == "00":
                continue
            dry = bq.query(table.select_sql, job_config=bigquery.QueryJobConfig(
                dry_run=True, use_query_cache=False, query_parameters=params, maximum_bytes_billed=1024**3,
            ))
            aliases = {"FLOAT": "FLOAT64", "INTEGER": "INT64", "BOOLEAN": "BOOL"}
            actual = [
                (field.name, f"ARRAY<{field.field_type}>" if field.mode == "REPEATED"
                 else aliases.get(field.field_type, field.field_type)) for field in dry.schema
            ]
            if actual != [(field.name, field.type) for field in table.columns]:
                raise ContractBlocked(f"{table.table_id}: SELECT output names/types differ from the BRD contract.")
        report["steps"].append({"step": phase, "actor": "Python/BigQuery dry-run", "status": "PASS"})
        phase = "GCS upload and 00 load"
        for table in contract.tables:
            if table.layer != "00":
                continue
            path, rows = prepared[table.table_id]
            prefix = table.gcs_prefix.replace("YYYY-MM-DD", now.date().isoformat())
            prefix = prefix.replace("YYYY", now.strftime("%Y")).replace("MM", now.strftime("%m")).replace("DD", now.strftime("%d"))
            name = path.with_suffix(".csv").name
            fields = [column for column in table.columns if not column.metadata]
            buffer = io.StringIO()
            writer = csv.DictWriter(buffer, fieldnames=[column.name for column in fields], lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
            data = path.read_bytes() if path.suffix.lower() == ".csv" else buffer.getvalue().encode()
            obj = ensure_object(bucket, prefix.rstrip("/") + "/" + name, data, content_type="text/csv")
            report["objects"].append(obj)
            report["jobs"].append(load_rows(
                bq, bucket, obj, table.table_id, table.table_id, table.columns,
                rows, now.date(), now, report["brd_sha256"],
                {column.name: column.metadata for column in table.columns if column.metadata},
            ))
        report["steps"].append({"step": phase, "actor": "Python/GCS generation and atomic load jobs", "status": "PASS"})
        phase = "atomic curated publication"
        job = bq.query(transform, job_config=bigquery.QueryJobConfig(
            query_parameters=params, maximum_bytes_billed=1024**3, labels=LABELS,
        ))
        job.result()
        report["jobs"].append({"phase": "01", "job_id": job.job_id})
        if kpi_transform:
            phase = "02 DSP staging, validation and atomic publication"
            job_02 = bq.query(kpi_transform, job_config=bigquery.QueryJobConfig(
                query_parameters=params, maximum_bytes_billed=1024**3, labels=LABELS,
            ))
            report["kpi_validation_failures"] = [dict(row) for row in job_02.result()]
            report["jobs"].append({"phase": "02", "job_id": job_02.job_id})
            if report["kpi_validation_failures"]:
                raise ContractBlocked("02 publication blocked; see keyed validation diagnostics.")
            report["steps"].append({"step": phase, "actor": "Python/BigQuery DSP CLI adaptation", "status": "PASS"})
        for table, target, _ in wanted:
            actual = bq.get_table(target)
            schema_actions(actual, target, allow_nullable_additions=False)
            report["row_counts"][table.table_id] = actual.num_rows
        report["steps"].append({"step": phase, "actor": "Python/BigQuery validated SQL", "status": "PASS"})
        if verify_retry:
            report["idempotency_verification"] = verify_unchanged_retry(
                bq, transform, params, sum(table.layer == "01" for table in contract.tables),
            )
            if kpi_transform:
                report["idempotency_verification_02"] = verify_unchanged_retry(
                    bq, kpi_transform, params, len(kpi_tables),
                )
        report["status"] = "SUCCESS"
        write_report(report, directory)
        for name, content_type in (("run_report.md", "text/markdown"), ("run_report.json", "application/json")):
            ensure_object(
                bucket, f"run_reports/{run_id}/attempts/{now:%Y%m%dT%H%M%S%f}/{name}",
                (directory / name).read_bytes(), content_type=content_type,
            )
    except (ValueError, LLMError, GoogleAPICallError, GoogleAuthError, OSError) as error:
        report["status"] = "BLOCKED" if isinstance(error, ContractBlocked) else "FAILED"
        reason = str(error) if type(error) in {ContractBlocked, ValueError} else type(error).__name__
        report["issues"].append({"check": phase, "result": reason,
                                 "action": "Correct the BRD/source/IAM issue and rerun; no implicit defaults or destructive fallback."})
        report["steps"].append({"step": phase, "actor": "Python validation gate", "status": report["status"]})
    finally:
        write_report(report, directory)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--brd", type=Path, required=True)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, default=PROJECT_ROOT / "generated")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--verify-idempotency", action="store_true")
    args = parser.parse_args()
    try:
        text = read_document(args.brd).text
    except (OSError, ValueError):
        # The generic runner persists the actual read/auth failure as a correction report.
        text = ""
    rainfall = all(name in text for name in (
        "01_dim_station_watershed", "01_fact_daily_station_rainfall",
        "01_fact_watershed_station_coverage", "01_fact_station_rainfall_assessment",
    ))
    runner = execute_rainfall if rainfall else execute_generic
    report = runner(args.brd, args.source_root, args.output_root, dry_run=args.dry_run,
                    verify_retry=args.verify_idempotency)
    print(json.dumps({key: report.get(key) for key in (
        "status", "run_id", "report_directory", "row_counts", "issues",
    )}, indent=2))
    return 0 if report["status"] in {"SUCCESS", "DRY_RUN_READY"} else 2


if __name__ == "__main__":
    raise SystemExit(main())
