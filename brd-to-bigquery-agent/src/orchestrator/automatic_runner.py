"""Non-interactive DEV provisioning for validated GCS -> 00 -> 01 contracts."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from google.api_core.exceptions import NotFound
from pydantic import BaseModel, ConfigDict, Field, StrictBool

from ..brd_parser.document_reader import read_document
from ..brd_parser.spec_validator import validate_spec
from ..file_prep.gcs_uploader import GcsUploader
from ..file_prep.xlsx_to_csv import normalize_header
from ..llm.extraction_agent import LLMExtractionAgent
from ..llm.provider import get_provider
from ..models.pipeline_spec import PipelineSpec
from ..models.pipeline_spec import BigQueryType, ColumnSpec
from ..sql_generator import generate_00_sql, generate_01_sql
from ..sql_generator.render_template import write_sql
from ..utils.config import AppConfig, PROJECT_ROOT


class ProposalColumn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str
    type: BigQueryType
    nullable: StrictBool = True
    preserve_leading_zeros: StrictBool = False
    description: str | None = None


class ProposalTable(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source_name: str
    table_00: str
    table_01: str
    columns: list[ProposalColumn] = Field(min_length=1)
    gcs_folder: str | None = None
    grain: str = "other"
    source_location_hint: str | None = None
    source_format: str | None = None
    partition_by: str | None = None
    cluster_by: list[str] | None = None


class ProposalDirectives(BaseModel):
    model_config = ConfigDict(extra="forbid")
    partition_by: str | None = None
    cluster_candidates: list[str] = Field(default_factory=list)
    load_date_field: str | None = None
    rerun_strategy: str | None = None


class AutomaticProposal(BaseModel):
    model_config = ConfigDict(extra="forbid")
    tables: list[ProposalTable] = Field(min_length=1)
    project_id: str | None = None
    gcs_bucket: str | None = None
    dataset_00: str | None = None
    dataset_01: str | None = None
    schedule: str | None = None
    historical_runs: list[str] = Field(default_factory=list)
    directives: ProposalDirectives = Field(default_factory=ProposalDirectives)
    assumptions: list[dict[str, Any]] = Field(default_factory=list)
    unsupported: list[str] = Field(default_factory=list)


class AutomaticExtractionAgent(LLMExtractionAgent):
    def _document_text(self, document) -> str:
        if len(document.text) > self.max_document_chars:
            raise ValueError("BRD exceeds the supported context size; refusing partial extraction.")
        return super()._document_text(document)

    def _compile(self, payload, document, run_id):
        if payload.get("assumptions") or payload.get("unsupported"):
            raise ValueError("BRD has unresolved assumptions or requirements outside the supported table contract.")
        payload = AutomaticProposal.model_validate(payload).model_dump(exclude_none=True)
        spec = super()._compile(payload, document, run_id)
        for raw, table in zip(payload["tables"], spec.tables):
            if table.table_00 not in document.text or table.table_01 not in document.text:
                raise ValueError("Physical 00 and 01 names must appear in the BRD; do not invent them.")
            if raw.get("partition_by"):
                table.partition_by = raw["partition_by"]
            if raw.get("cluster_by") is not None:
                table.cluster_by = raw["cluster_by"]
        if self.config.bigquery.dataset_01 is None and spec.dataset_01 not in document.text:
            raise ValueError("The 01 dataset must be specified in the BRD or pinned by environment.")
        return spec

    def _compile_column(self, raw):
        if not raw.get("type"):
            raise ValueError("Every business field requires an explicit BigQuery type.")
        column = super()._compile_column(raw)
        if "nullable" in raw:
            if not isinstance(raw["nullable"], bool):
                raise ValueError("Column nullable must be a boolean.")
            column.nullable = raw["nullable"]
        return column


def environment_config(environment: dict[str, str]) -> AppConfig:
    def value(name: str) -> str:
        result = environment.get(name, "").strip()
        if not result or result.startswith("<"):
            raise ValueError(f"Configure {name} before automatic provisioning.")
        return result

    project = value("GOOGLE_CLOUD_PROJECT")
    if not re.fullmatch(r"[a-z][a-z0-9-]{4,28}[a-z0-9]", project):
        raise ValueError("GOOGLE_CLOUD_PROJECT is not a valid project identifier.")
    return AppConfig(
        environment="dev",
        gcp={"project_id": project, "location": value("GOOGLE_CLOUD_LOCATION")},
        gcs={"bucket": value("GCS_LANDING_BUCKET"), "overwrite_policy": "fail"},
        bigquery={
            "dataset_00": value("BQ_DATASET_00"),
            "dataset_01": (
                environment.get("BQ_DATASET_01")
                if environment.get("BQ_DATASET_01") and not environment["BQ_DATASET_01"].startswith("<")
                else None
            ),
        },
        llm={"enabled": True, "provider": value("LLM_PROVIDER"), "model": value("LLM_MODEL")},
    )


def validate_value(text: str, column: ColumnSpec) -> None:
    if not text:
        if not column.nullable:
            raise ValueError(f"Blank required field {column.name}.")
        return
    try:
        if column.type in {"INT64", "INTEGER"}:
            if not re.fullmatch(r"[+-]?\d+", text):
                raise ValueError()
        elif column.type in {"FLOAT64", "NUMERIC", "BIGNUMERIC"}:
            if not Decimal(text).is_finite():
                raise ValueError()
        elif column.type == "DATE":
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
                raise ValueError()
            date.fromisoformat(text)
        elif column.type in {"TIMESTAMP", "DATETIME"}:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            if column.type == "TIMESTAMP" and parsed.tzinfo is None:
                raise ValueError()
        elif column.type == "BOOL" and text.lower() not in {"true", "false"}:
            raise ValueError()
    except (ValueError, InvalidOperation):
        raise ValueError(f"Invalid {column.type} value for {column.name}; record rejected.") from None


def prepare_sources(spec: PipelineSpec, source_root: Path, output: Path) -> dict[str, list[Path]]:
    """Match complete business headers, rejecting ambiguous sheets and files."""
    from openpyxl import load_workbook

    candidates: list[tuple[Path, str | None, list[str], list[list[Any]]]] = []
    for path in sorted(source_root.iterdir()):
        if path.name.startswith("~$") or not path.is_file():
            continue
        if path.suffix.lower() == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as handle:
                records = list(csv.reader(handle))
            if records:
                candidates.append((path, None, records[0], records[1:]))
        elif path.suffix.lower() in {".xlsx", ".xlsm"}:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    records = list(sheet.iter_rows(values_only=True))
                    if records:
                        candidates.append((
                            path, sheet.title,
                            [str(cell or "").strip() for cell in records[0]],
                            [list(row) for row in records[1:] if any(cell is not None for cell in row)],
                        ))
            finally:
                workbook.close()

    result: dict[str, list[Path]] = {}
    used: set[tuple[Path, str | None]] = set()
    for table in spec.tables:
        wanted = [column.name for column in table.file_columns]
        normalized = [normalize_header(name) for name in wanted]
        matches = [
            item for item in candidates
            if len(item[2]) == len(wanted)
            and len(set(normalize_header(header) for header in item[2])) == len(wanted)
            and set(normalize_header(header) for header in item[2]) == set(normalized)
        ]
        if not matches:
            raise ValueError(f"{table.table_00}: no file has the complete BRD business schema.")
        paths = [item[0] for item in matches]
        if len(paths) != len(set(paths)):
            raise ValueError(f"{table.table_00}: multiple matching workbook sheets need an explicit mapping.")
        outputs = []
        for path, sheet, headers, rows in matches:
            if (path, sheet) in used:
                raise ValueError(f"{path.name}: ambiguous match to multiple feeds.")
            used.add((path, sheet))
            if not rows:
                raise ValueError(f"{path.name}: no data rows.")
            order = [list(map(normalize_header, headers)).index(name) for name in normalized]
            staged = output / table.gcs_folder / table.csv_file_name(
                spec.run_id, hashlib.sha256(path.name.encode()).hexdigest()[:12]
            )
            staged.parent.mkdir(parents=True, exist_ok=True)
            with staged.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle, quoting=csv.QUOTE_NONE, escapechar="\\")
                writer.writerow(table.csv_columns)
                for row in rows:
                    if len(row) != len(headers):
                        raise ValueError(f"{path.name}: jagged record.")
                    values = {}
                    for column, index in zip(table.file_columns, order):
                        cell = row[index]
                        if path.suffix.lower() != ".csv" and column.type == "STRING" and isinstance(cell, (int, float)):
                            raise ValueError(f"{path.name}: numeric Excel identifier needs reviewed string formatting.")
                        text = "" if cell is None else str(cell)
                        validate_value(text, column)
                        if any(char in text for char in (",", "\n", "\r")):
                            raise ValueError(f"{path.name}: value incompatible with standard unquoted CSV template.")
                        values[column.name] = text
                    values["run_date"] = spec.run_date.isoformat()
                    values["Source_FileName"] = path.name
                    writer.writerow([values[column.name] for column in table.columns_00])
            outputs.append(staged)
        result[table.source_name] = outputs
    return result


def provision(
    spec: PipelineSpec, config: AppConfig, files: dict[str, list[Path]], sql: dict[str, dict[str, str]],
    audit: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Preflight all targets before uploads; never replace tables or objects."""
    from google.cloud import bigquery

    client = bigquery.Client(project=spec.project_id, location=spec.location)
    uploader = GcsUploader(config, dry_run=False)
    bucket = uploader.client().get_bucket(spec.gcs_bucket)
    if bucket.location.upper() != spec.location.upper():
        raise ValueError("Bucket and dataset locations differ.")
    for dataset_id in (spec.dataset_00, spec.dataset_01):
        try:
            existing = client.get_dataset(f"{spec.project_id}.{dataset_id}")
        except NotFound:
            continue
        if existing.location.upper() != spec.location.upper():
            raise ValueError("Existing dataset location differs from the configured location.")
    for table in spec.tables:
        for dataset, name in ((spec.dataset_00, table.table_00), (spec.dataset_01, table.table_01)):
            try:
                client.get_table(f"{spec.project_id}.{dataset}.{name}")
            except NotFound:
                continue
            raise ValueError(f"{name}: existing table requires a schema/rerun plan; automatic replacement is prohibited.")
        for path in files[table.source_name]:
            if bucket.blob(f"{table.gcs_folder}/{path.name}").exists():
                raise ValueError("An upload target already exists; automatic overwrite is prohibited.")
    audit = audit if audit is not None else []
    for table in spec.tables:
        for path in files[table.source_name]:
            blob = bucket.blob(f"{table.gcs_folder}/{path.name}")
            blob.upload_from_filename(str(path), content_type="text/csv", if_generation_match=0)
            audit.append({"action": "upload", "object": blob.name, "generation": blob.generation})
    for dataset_id in (spec.dataset_00, spec.dataset_01):
        dataset = bigquery.Dataset(f"{spec.project_id}.{dataset_id}")
        dataset.location = spec.location
        client.create_dataset(dataset, exists_ok=True)
        audit.append({"action": "ensure_dataset", "dataset": dataset_id})
    for table in spec.tables:
        for layer in ("00", "01"):
            job = client.query(sql[table.source_name][layer], job_config=bigquery.QueryJobConfig(
                labels={"framework": "brd_to_bigquery_agent", "step": f"automatic_{layer}"}
            ))
            job.result()
            target = table.table_00 if layer == "00" else table.table_01
            dataset = spec.dataset_00 if layer == "00" else spec.dataset_01
            client.get_table(f"{spec.project_id}.{dataset}.{target}")
            audit.append({"action": f"create_{layer}", "table": target, "job_id": job.job_id})
    return audit


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--brd-path", type=Path, required=True)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--run-id", required=True, help="YYYY_MM, following existing framework conventions.")
    parser.add_argument("--dry-run", action="store_true", help="Generate local plans without cloud writes.")
    args = parser.parse_args(argv)
    identity = hashlib.sha256(
        (str(args.brd_path.resolve()) + args.run_id).encode()
    ).hexdigest()[:16]
    report: dict[str, Any] = {
        "status": "STARTED", "dry_run": args.dry_run, "cloud_operations": [],
        "brd_file": args.brd_path.name, "run_id": args.run_id,
    }
    report_path = PROJECT_ROOT / f"generated/reports/automatic_run_{identity}.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        config = environment_config(dict(os.environ))
        document = read_document(args.brd_path)
        provider = get_provider(config.llm.provider, config.llm.model)
        result = AutomaticExtractionAgent(provider, config, max_attempts=2).extract(document, args.run_id)
        report["extraction"] = result.to_dict()
        if not result.ok or result.spec is None or result.unsupported or result.assumptions:
            raise ValueError("BRD extraction is incomplete or has unresolved assumptions/unsupported requirements.")
        spec = result.spec
        validate_spec(spec, config).raise_if_failed("Automatic specification")
        for table in spec.tables:
            for identifier in (table.table_00, table.table_01, table.gcs_folder, table.file_pattern):
                if not re.fullmatch(r"[A-Za-z0-9_]+", identifier):
                    raise ValueError("Unsafe table or object identifier.")
            if table.table_00 not in document.text or table.table_01 not in document.text:
                raise ValueError("Both physical table names must be explicitly grounded in the BRD.")
            if any(column.name not in document.text for column in table.file_columns):
                raise ValueError("Business column is absent from the BRD.")
            for column in table.schema_columns:
                if column.description and any(char in column.description for char in ("\n", "\r")):
                    raise ValueError("Column descriptions cannot contain SQL line breaks.")
            partition = next(c for c in table.schema_columns if c.name == table.partition_by)
            if partition.type != "DATE":
                raise ValueError("The current automatic template supports DATE partitioning only.")
            if len(table.cluster_by) > 4 or any(
                c.type in {"FLOAT64", "BYTES"} for c in table.schema_columns if c.name in table.cluster_by
            ):
                raise ValueError("Unsupported clustering contract.")
            table.gcs_folder = f"agentic_brd/{hashlib.sha256(args.brd_path.read_bytes()).hexdigest()[:16]}/{table.table_00}"
        spec.save(PROJECT_ROOT / "generated/specs")
        files = prepare_sources(spec, args.source_root, PROJECT_ROOT / "generated/staging_csv" / spec.run_id)
        sql = {}
        for table in spec.tables:
            sql[table.source_name] = {
                "00": generate_00_sql(spec, table, config).replace(
                    "CREATE OR REPLACE EXTERNAL TABLE", "CREATE EXTERNAL TABLE IF NOT EXISTS"
                ),
                "01": generate_01_sql(spec, table, config),
            }
            for layer, statement in sql[table.source_name].items():
                write_sql(statement, PROJECT_ROOT / "generated/sql" / layer,
                          f"{table.table_00 if layer == '00' else table.table_01}.sql")
        report["tables"] = [{"00": t.table_00, "01": t.table_01} for t in spec.tables]
        if not args.dry_run:
            provision(spec, config, files, sql, audit=report["cloud_operations"])
        report["status"] = "PLANNED" if args.dry_run else "PROVISIONED"
    except Exception as exc:
        report["status"] = "FAILED"
        report["error_type"] = type(exc).__name__
        if isinstance(exc, ValueError):
            report["reason"] = str(exc)
        print(f"Automatic provisioning stopped ({type(exc).__name__}); run report: {report_path}")
        return 2
    finally:
        report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"Automatic provisioning: {report['status']}")
    print(f"Run report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
