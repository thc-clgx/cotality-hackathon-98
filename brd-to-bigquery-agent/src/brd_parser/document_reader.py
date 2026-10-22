"""Reads a BRD into a flat, indented line model.

Supported formats: ``.docx``, text-based ``.pdf``, ``.txt``, ``.md``.

Indent level matters: a BRD is essentially a nested bullet list, and the
extractor relies on "these Fields bullets belong to the Brushfire heading".
For ``.docx`` the level comes from the paragraph indent / list level; for text
files it comes from leading whitespace or bullet depth.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from ..utils.file_access import cleanup_copy, local_readable_copy
from ..utils.logging_setup import get_logger

logger = get_logger(__name__)

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".txt", ".md", ".markdown"}
_BULLET = re.compile(r"^\s*(?:[-*\u2022\u25cf\u25aa\u2023o]|\d+[.)]|[a-z][.)])\s+", re.IGNORECASE)


@dataclass
class BrdLine:
    """One logical line of the BRD."""

    index: int
    text: str
    level: int = 0
    style: str = "Normal"
    source: str = "paragraph"

    @property
    def normalized(self) -> str:
        return re.sub(r"\s+", " ", self.text).strip().lower()


@dataclass
class BrdDocument:
    path: Path
    lines: list[BrdLine] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(line.text for line in self.lines)

    def find(self, pattern: str, *, regex: bool = False, start: int = 0) -> int:
        """Index of the first line matching ``pattern`` (case-insensitive), else -1."""
        for line in self.lines[start:]:
            haystack = line.normalized
            if regex:
                if re.search(pattern, haystack, re.IGNORECASE):
                    return line.index
            elif pattern.lower() in haystack:
                return line.index
        return -1

    def section(self, start: int, *, max_level: int | None = None, stop_before: int | None = None) -> list[BrdLine]:
        """Lines nested under ``lines[start]`` (exclusive of the heading itself)."""
        if start < 0 or start >= len(self.lines):
            return []
        base_level = self.lines[start].level if max_level is None else max_level
        collected: list[BrdLine] = []
        for line in self.lines[start + 1:]:
            if stop_before is not None and line.index >= stop_before:
                break
            if line.level <= base_level:
                break
            collected.append(line)
        return collected


def _read_docx(path: Path) -> BrdDocument:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ImportError("Reading .docx requires python-docx. pip install python-docx") from exc

    document = docx.Document(str(path))
    doc = BrdDocument(path=path)
    index = 0

    for paragraph in document.paragraphs:
        raw = (paragraph.text or "").strip()
        if not raw:
            continue
        level = _docx_level(paragraph)
        # Word sometimes packs several bullets into a single paragraph.
        for chunk in raw.split("\n"):
            chunk = chunk.strip()
            if not chunk:
                continue
            doc.lines.append(
                BrdLine(index=index, text=chunk, level=level, style=paragraph.style.name if paragraph.style else "Normal")
            )
            index += 1

    for table in document.tables:
        rows = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
        doc.tables.append(rows)
        for row in rows:
            joined = " | ".join(c for c in row if c)
            if joined:
                doc.lines.append(BrdLine(index=index, text=joined, level=0, style="Table", source="table"))
                index += 1

    return doc


def _docx_level(paragraph) -> int:
    """Approximate the outline level of a Word paragraph.

    Word bullets carry their depth in ``w:numPr/w:ilvl``; ``left_indent`` is only
    a fallback because manual indents are noisy.
    """
    try:
        ilvl = paragraph._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
        if ilvl:
            return int(ilvl[0])
    except (AttributeError, ValueError, TypeError):
        pass

    style_name = (paragraph.style.name if paragraph.style else "") or ""
    match = re.search(r"(\d+)$", style_name)
    if style_name.startswith("Heading") and match:
        return max(0, int(match.group(1)) - 1)

    indent = paragraph.paragraph_format.left_indent
    if indent is not None:
        # Word list levels sit roughly one 0.25" (228600 EMU) step apart.
        return max(0, min(9, int(round(indent.emu / 228600))))
    return 0


def _read_text(path: Path) -> BrdDocument:
    doc = BrdDocument(path=path)
    index = 0
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if not raw.strip():
            continue
        expanded = raw.replace("\t", "    ")
        leading = len(expanded) - len(expanded.lstrip(" "))
        text = _BULLET.sub("", expanded).strip()
        text = re.sub(r"^#+\s*", "", text)
        heading_level = 0
        if raw.lstrip().startswith("#"):
            heading_level = len(raw.lstrip()) - len(raw.lstrip().lstrip("#")) - 1
        level = heading_level if heading_level else leading // 2
        doc.lines.append(BrdLine(index=index, text=text, level=level, style="Text"))
        index += 1
    return doc


def _read_pdf(path: Path) -> BrdDocument:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError("Reading .pdf requires pypdf; install the framework requirements.") from None

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("Encrypted BRDs require an unencrypted, approved local copy.")
    if not reader.pages:
        raise ValueError("The PDF BRD has no pages.")
    doc = BrdDocument(path=path)
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text(extraction_mode="layout") if "/Contents" in page else ""
        text = text or ""
        if not text.strip():
            raise ValueError(
                f"PDF page {page_number} has no extractable text; supply an OCR-reviewed "
                "text-based PDF or DOCX. Partial BRD extraction is not permitted."
            )
        for raw in text.splitlines():
            if raw.strip():
                doc.lines.append(
                    BrdLine(
                        index=len(doc.lines),
                        text=raw.strip(),
                        style=f"PDF page {page_number}",
                        source=f"page:{page_number}",
                    )
                )
    return doc


def read_document(path: str | Path) -> BrdDocument:
    """Load a BRD from disk into a :class:`BrdDocument`."""
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"BRD not found: {source}")

    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported BRD format {suffix!r}. Supported: {sorted(SUPPORTED_SUFFIXES)}")

    readable = local_readable_copy(source)
    try:
        if suffix == ".docx":
            doc = _read_docx(readable)
        elif suffix == ".pdf":
            doc = _read_pdf(readable)
        else:
            doc = _read_text(readable)
    finally:
        cleanup_copy(source, readable)
    doc.path = source
    logger.info("Read BRD %s (%s lines, %s tables)", source.name, len(doc.lines), len(doc.tables))
    return doc
