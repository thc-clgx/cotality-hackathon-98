"""Offline synthetic regression tests for ingestion and column profiling."""

from __future__ import annotations

import io
import unittest

from brd_insight_agent.loader import correct_column


class CorrectionTests(unittest.TestCase):
    def test_exact_numeric_and_date_formats(self):
        import pandas as pd
        frame = pd.DataFrame({"Amount": ["1.234,56", "2.345,67", None], "Date": ["31/12/2025", "01/02/2026", None]})
        numeric = correct_column(frame, "Amount", "numeric", decimal=",", thousands=".")
        self.assertEqual(numeric["Amount"].dropna().tolist(), [1234.56, 2345.67])
        dated = correct_column(frame, "Date", "datetime", date_format="%d/%m/%Y")
        self.assertEqual(str(dated["Date"].iloc[0].date()), "2025-12-31")
        self.assertEqual(frame["Amount"].iloc[0], "1.234,56")

    def test_bad_values_reject_whole_correction(self):
        import pandas as pd
        frame = pd.DataFrame({"Amount": ["1", "bad"]})
        with self.assertRaises(ValueError):
            correct_column(frame, "Amount", "numeric")
        self.assertEqual(frame["Amount"].tolist(), ["1", "bad"])
from unittest.mock import patch

import openpyxl
import pandas as pd
from pandas.api import types as ptypes

from brd_insight_agent.loader import (
    clean_frame,
    load_sheet,
    load_workbook,
    workbook_sheet_names,
)
from brd_insight_agent.profiler import ColumnProfile, _tokens, by_role, is_rate_like, profile_frame


def workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class CleaningTests(unittest.TestCase):
    def test_all_200001_numeric_rows_are_retained(self):
        source = pd.DataFrame({"Revenue": range(200_001)})
        result = clean_frame(source)
        pd.testing.assert_frame_equal(result, source)
        self.assertEqual(result["Revenue"].iloc[-1], 200_000)

    def test_csv_200001_rows_and_last_minority_text_are_retained(self):
        content = b"value\n" + b"1\n" * 200_000 + b"not numeric\n"
        result = clean_frame(load_workbook(content, "large.csv")["csv"])
        self.assertEqual(len(result), 200_001)
        self.assertEqual(result["value"].iloc[0], "1")
        self.assertEqual(result["value"].iloc[-1], "not numeric")
        self.assertEqual(int(result["value"].isna().sum()), 0)

    def test_blanks_are_normalized_before_empty_rows_and_columns_are_dropped(self):
        source = pd.DataFrame({
            "amount": ["1", " \t ", None, "2", ""],
            "label": ["a", "", None, "b", " \t "],
            "empty": [" ", "\n", None, pd.NA, ""],
        })
        original = source.copy(deep=True)
        result = clean_frame(source)
        self.assertEqual(list(result.columns), ["amount", "label"])
        self.assertEqual(result["amount"].tolist(), [1, 2])
        self.assertEqual(result["label"].tolist(), ["a", "b"])
        self.assertEqual(result.index.tolist(), [0, 1])
        pd.testing.assert_frame_equal(source, original)

    def test_numeric_coercion_with_nullable_and_inferred_string_dtypes(self):
        for dtype in (object, "string", pd.StringDtype(storage="python"), None):
            with self.subTest(dtype=dtype):
                values = pd.Series([" 1,234.50 ", "-2.5", "", " \t", None, pd.NA], dtype=dtype)
                result = clean_frame(pd.DataFrame({"value": values, "anchor": range(6)}))
                self.assertTrue(ptypes.is_numeric_dtype(result["value"].dtype))
                self.assertEqual(result["value"].iloc[:2].tolist(), [1234.5, -2.5])
                self.assertTrue(result["value"].iloc[2:].isna().all())

    def test_blank_normalization_also_applies_to_unconverted_text_and_categories(self):
        for dtype in (object, "string", "category", None):
            with self.subTest(dtype=dtype):
                source = pd.Series([" keep spaces ", "", " \n\t ", None, pd.NA], dtype=dtype)
                result = clean_frame(pd.DataFrame({"text": source, "anchor": range(5)}))
                self.assertEqual(result["text"].iloc[0], " keep spaces ")
                self.assertTrue(result["text"].iloc[1:].isna().all())

    def test_duplicate_normalized_headers_reserve_existing_suffixes(self):
        source = pd.DataFrame([[1, 2, 3, 4, 5]], columns=[" a ", "a", "a_1", "a", "a_2"])
        result = clean_frame(source)
        self.assertEqual(result.columns.tolist(), ["a", "a_3", "a_1", "a_4", "a_2"])
        self.assertEqual(result.iloc[0].tolist(), [1, 2, 3, 4, 5])
        basic = clean_frame(pd.DataFrame([[1, 2, 3]], columns=["a", "a", "a_1"]))
        self.assertEqual(basic.columns.tolist(), ["a", "a_2", "a_1"])

    def test_empty_and_generated_headers_cannot_collide(self):
        source = pd.DataFrame(
            [[1] * 8],
            columns=[None, "column_1", "", "Unnamed: 3", "column_4", float("nan"), 7, " 7 "],
        )
        result = clean_frame(source)
        self.assertTrue(result.columns.is_unique)
        self.assertTrue(all(isinstance(name, str) and name.strip() for name in result.columns))
        self.assertEqual(result.shape, source.shape)
        pd.testing.assert_frame_equal(clean_frame(result), result)

    def test_ambiguous_columns_preserve_every_nonblank_original_value(self):
        cases = [
            ["12"] * 9 + [" minority text "],
            ["2026-09-01"] * 9 + ["not a date"],
            ["2026-09-01", "2026-02-30"],
            ["2026-09-01", "42"],
            ["2026-09-01", 42],
            ["03/04/2026", "04/05/2026"],
            ["January 2026", "February 2026"],
            ["50%", "0.5"],
            ["50%", 0.5],
            ["50%", "50"],
            ["50%", "$50"],
            ["$10", "\u20ac20"],
            ["$10", "20"],
            ["$10", 20],
            ["\u20ac1.234,56", "\u20ac2.345,67"],
            ["1,23", "4,56"],
            ["1 2", "34"],
            ["1e400", "2"],
            ["2026-09-01T00:00:00Z", "2026-09-02T00:00:00"],
            ["2026-09-01T00:00:00+01:00", "2026-09-02T00:00:00+02:00"],
        ]
        for values in cases:
            with self.subTest(values=values):
                source = pd.Series(values + [" \t ", None], name="value", dtype=object)
                result = clean_frame(pd.DataFrame({"value": source, "anchor": range(len(source))}))
                self.assertEqual(result["value"].iloc[:len(values)].tolist(), values)
                self.assertTrue(result["value"].iloc[len(values):].isna().all())

    def test_percent_only_columns_become_fractions(self):
        for dtype in (object, "string", None):
            with self.subTest(dtype=dtype):
                values = pd.Series(["50%", " 12.5 % ", "-2%", "100%", "", None], dtype=dtype)
                result = clean_frame(pd.DataFrame({"Rate": values, "anchor": range(6)}))
                self.assertTrue(ptypes.is_numeric_dtype(result["Rate"].dtype))
                self.assertEqual(result["Rate"].iloc[:4].tolist(), [0.5, 0.125, -0.02, 1.0])
                self.assertTrue(result["Rate"].iloc[4:].isna().all())

    def test_single_explicit_currency_parses_without_changing_units(self):
        for symbol in ("$", "\u20ac", "\u00a3"):
            with self.subTest(symbol=symbol):
                result = clean_frame(pd.DataFrame({
                    "Revenue": [f"{symbol}1,234.50", f"-{symbol}2.5", f"{symbol}-3.5", ""],
                    "anchor": range(4),
                }))
                self.assertEqual(result["Revenue"].iloc[:3].tolist(), [1234.5, -2.5, -3.5])
                self.assertTrue(pd.isna(result["Revenue"].iloc[3]))

    def test_leading_zero_values_and_named_ids_are_not_coerced(self):
        cases = [
            ("value", ["001", "012", "42"]),
            ("value", ["-001", "+012", "42"]),
            ("CustomerID", ["001", "012", "42"]),
            ("OrderNumber", ["123", "456", "789"]),
            ("AccountCode", ["2026-09-01", "2026-09-02", "2026-09-03"]),
        ]
        for name, values in cases:
            with self.subTest(name=name, values=values):
                result = clean_frame(pd.DataFrame({name: values}))
                self.assertEqual(result[name].tolist(), values)

    def test_only_plausible_explicit_iso_dates_are_inferred(self):
        source = pd.DataFrame({
            "timestamp": ["2026-09-01", "2026-09-02T13:45:00.123", "", None],
            "anchor": range(4),
        })
        result = clean_frame(source)
        self.assertTrue(ptypes.is_datetime64_any_dtype(result["timestamp"].dtype))
        self.assertEqual(result["timestamp"].iloc[0], pd.Timestamp("2026-09-01"))
        self.assertEqual(result["timestamp"].iloc[1], pd.Timestamp("2026-09-02T13:45:00.123"))
        self.assertTrue(result["timestamp"].iloc[2:].isna().all())
        for values in (["20260901", "20260902"], ["1690000000", "1690000001"]):
            compact = clean_frame(pd.DataFrame({"value": values}))
            self.assertFalse(ptypes.is_datetime64_any_dtype(compact["value"].dtype))

    def test_uniform_iso_timezone_is_retained(self):
        result = clean_frame(pd.DataFrame({"timestamp": ["2026-09-01T12:00:00Z", "2026-09-02T12:00:00Z"]}))
        self.assertTrue(ptypes.is_datetime64_any_dtype(result["timestamp"].dtype))
        self.assertEqual(result["timestamp"].iloc[0], pd.Timestamp("2026-09-01T12:00:00Z"))


class ProfileTests(unittest.TestCase):
    def test_uniqueness_does_not_make_measures_dates_or_free_text_identifiers(self):
        source = pd.DataFrame({
            "Revenue": [i + 0.125 for i in range(40)],
            "timestamp": pd.date_range("2026-01-01", periods=40),
            "Description": [f"description {i}" for i in range(40)],
        })
        profiles = profile_frame(source)
        self.assertEqual(profiles["Revenue"].role, "numeric")
        self.assertEqual(profiles["timestamp"].role, "datetime")
        self.assertEqual(profiles["Description"].role, "freetext")
        self.assertTrue(all(not p.is_identifier for p in profiles.values()))

    def test_numeric_and_string_id_names_are_labels_even_when_repeated(self):
        profiles = profile_frame(pd.DataFrame({
            "CustomerID": [i % 3 for i in range(40)],
            "SKU": range(40),
            "InvoiceNumber": [i % 2 for i in range(40)],
            "OrderNo": [i % 2 for i in range(40)],
            "order_number_1": [i % 2 for i in range(40)],
            "CustomerCode": [f"C{i % 3}" for i in range(40)],
            "record_uuid": [f"record-{i}" for i in range(40)],
        }))
        self.assertTrue(all(p.is_identifier for p in profiles.values()))
        self.assertTrue(all(p.role != "numeric" for p in profiles.values()))
        self.assertEqual(profiles["CustomerID"].role, "categorical")
        self.assertEqual(profiles["SKU"].role, "freetext")
        self.assertTrue(profiles["SKU"].is_identifier)
        self.assertEqual(profiles["CustomerCode"].role, "categorical")
        self.assertEqual(profiles["record_uuid"].role, "freetext")

    def test_identifier_tokens_do_not_match_substrings_or_count_names(self):
        names = ["Width", "Paid", "Liquidity", "RenumberedTotal", "UnitsNumbered", "NumberOfOrders", "number_of_units", "number"]
        profiles = profile_frame(pd.DataFrame({name: range(40) for name in names}))
        self.assertTrue(all(p.role == "numeric" and not p.is_identifier for p in profiles.values()))

    def test_profile_helpers_and_rate_heuristics_remain_compatible(self):
        for name in ("DiscountPct", "MarginPct", "AverageRevenue", "Score", "CostPerUnit"):
            self.assertTrue(is_rate_like(name), name)
        for name in ("Revenue", "CorporateRevenue", "SeparateCost"):
            self.assertFalse(is_rate_like(name), name)
        self.assertEqual(_tokens("DiscountPct"), {"discount", "pct"})
        profile = ColumnProfile("Revenue", "numeric", 3, 0.0, False)
        self.assertEqual(by_role({"Revenue": profile}, "numeric"), [profile])
        self.assertEqual(by_role({"Revenue": profile}, "datetime"), [])


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.content = workbook_bytes({
            "Plain": [["OrderID", "Revenue"], ["0001", 2.5], ["0002", 3.5]],
            "Report": [["Sales report"], [], ["CustomerCode", "Rate"], ["007", "25%"], ["008", "50%"]],
        })

    def test_sheet_names_do_not_parse_any_sheet(self):
        with patch.object(pd.ExcelFile, "parse", side_effect=AssertionError("sheet data must not be parsed")):
            self.assertEqual(workbook_sheet_names(self.content, "book.xlsx"), ["Plain", "Report"])
        self.assertEqual(workbook_sheet_names(b"a,b\n1,2\n", "book.CSV"), ["csv"])

    def test_only_selected_sheet_is_parsed_with_its_own_header(self):
        with patch.object(pd.ExcelFile, "parse", autospec=True, side_effect=pd.ExcelFile.parse) as parse:
            report = clean_frame(load_sheet(self.content, "book.xlsx", "Report", header_row=2))
            self.assertEqual(parse.call_count, 1)
            self.assertEqual(parse.call_args.kwargs["sheet_name"], "Report")
        self.assertEqual(report.columns.tolist(), ["CustomerCode", "Rate"])
        self.assertEqual(report["CustomerCode"].tolist(), ["007", "008"])
        self.assertEqual(report["Rate"].tolist(), [0.25, 0.5])
        plain = clean_frame(load_sheet(self.content, "book.xlsx", "Plain"))
        self.assertEqual(plain["OrderID"].tolist(), ["0001", "0002"])
        self.assertEqual(plain["Revenue"].tolist(), [2.5, 3.5])

    def test_load_workbook_still_returns_every_sheet_and_accepts_header_row(self):
        sheets = load_workbook(self.content, "book.xlsx")
        self.assertEqual(list(sheets), ["Plain", "Report"])
        for name, frame in sheets.items():
            pd.testing.assert_frame_equal(frame, load_sheet(self.content, "book.xlsx", name))
        offset = workbook_bytes({"Report": [["Title"], [], ["Code", "Amount"], ["001", 5]]})
        sheets = load_workbook(offset, "offset.xlsx", 2)
        self.assertEqual(sheets["Report"].columns.tolist(), ["Code", "Amount"])
        self.assertEqual(sheets["Report"].iloc[0].tolist(), ["001", 5])

    def test_csv_header_selection_preserves_ids_and_literal_na_labels(self):
        content = b"Report\nAccountID,Region,Revenue\n0001,NA,12.5\n0002,NULL,25\n"
        result = clean_frame(load_sheet(content, "report.csv", "csv", header_row=1))
        self.assertEqual(result["AccountID"].tolist(), ["0001", "0002"])
        self.assertEqual(result["Region"].tolist(), ["NA", "NULL"])
        self.assertEqual(result["Revenue"].tolist(), [12.5, 25])
        pd.testing.assert_frame_equal(load_workbook(content, "report.csv", 1)["csv"], load_sheet(content, "report.csv", "csv", 1))

    def test_native_excel_types_remain_compatible_without_parsing_text_ids(self):
        content = workbook_bytes({"Sheet": [
            ["Active", "timestamp", "Revenue", "AccountID"],
            [True, pd.Timestamp("2026-09-01"), 1.5, "001"],
            [False, pd.Timestamp("2026-09-02"), 2.5, "002"],
        ]})
        result = load_sheet(content, "native.xlsx", "Sheet")
        self.assertTrue(ptypes.is_bool_dtype(result["Active"].dtype))
        self.assertTrue(ptypes.is_datetime64_any_dtype(result["timestamp"].dtype))
        self.assertTrue(ptypes.is_numeric_dtype(result["Revenue"].dtype))
        self.assertEqual(result["AccountID"].tolist(), ["001", "002"])
        profiles = profile_frame(clean_frame(result))
        self.assertEqual(profiles["Active"].role, "boolean")
        self.assertEqual(profiles["timestamp"].role, "datetime")
        self.assertFalse(profiles["Revenue"].is_identifier)

    def test_duplicate_headers_are_not_lost_in_csv_or_excel_parsing(self):
        contents = [
            (b"a,a,a_1\n1,2,3\n", "duplicate.csv", "csv"),
            (workbook_bytes({"Sheet": [["a", "a", "a_1"], [1, 2, 3]]}), "duplicate.xlsx", "Sheet"),
        ]
        for content, filename, sheet in contents:
            with self.subTest(filename=filename):
                result = clean_frame(load_sheet(content, filename, sheet))
                self.assertEqual(result.columns.tolist(), ["a", "a_2", "a_1"])
                self.assertEqual(result.iloc[0].tolist(), [1, 2, 3])

    def test_invalid_header_indices_fail_clearly(self):
        for header in (-1, 0.5, "0", None, True, [0]):
            for filename, content, sheet in (("data.csv", b"a\n1\n", "csv"), ("book.xlsx", self.content, "Plain")):
                with self.subTest(header=header, filename=filename):
                    with self.assertRaisesRegex(ValueError, "header_row.*non-negative integer"):
                        load_sheet(content, filename, sheet, header)
                    with self.assertRaisesRegex(ValueError, "header_row.*non-negative integer"):
                        load_workbook(content, filename, header)
        for filename, content, sheet in (("data.csv", b"a\n1\n", "csv"), ("book.xlsx", self.content, "Plain")):
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(ValueError, "header_row.*outside"):
                    load_sheet(content, filename, sheet, 100)
                with self.assertRaisesRegex(ValueError, "header_row.*outside"):
                    load_workbook(content, filename, 100)

    def test_unknown_sheets_fail_clearly(self):
        for filename, content in (("data.csv", b"a\n1\n"), ("book.xlsx", self.content)):
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(ValueError, "Unknown sheet"):
                    load_sheet(content, filename, "missing")

    def test_unsupported_extensions_are_rejected_before_parsing(self):
        for filename in ("data.txt", "data.ods", "data.xlsb", "data.xlsx.exe", "no_extension"):
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(ValueError, "Unsupported file extension"):
                    workbook_sheet_names(b"not parsed", filename)
                with self.assertRaisesRegex(ValueError, "Unsupported file extension"):
                    load_sheet(b"not parsed", filename, "Sheet")
                with self.assertRaisesRegex(ValueError, "Unsupported file extension"):
                    load_workbook(b"not parsed", filename)

    def test_empty_csv_fails_clearly_but_header_only_csv_is_valid(self):
        for content in (b"", b" \t\r\n", b"\xef\xbb\xbf", b"\xef\xbb\xbf \r\n"):
            with self.subTest(content=content):
                with self.assertRaisesRegex(ValueError, "CSV file is empty"):
                    workbook_sheet_names(content, "empty.csv")
                with self.assertRaisesRegex(ValueError, "CSV file is empty"):
                    load_sheet(content, "empty.csv", "csv")
                with self.assertRaisesRegex(ValueError, "CSV file is empty"):
                    load_workbook(content, "empty.csv")
        result = load_sheet(b"a,b\n", "header.csv", "csv")
        self.assertEqual(result.shape, (0, 2))
        self.assertEqual(result.columns.tolist(), ["a", "b"])

    def test_empty_excel_sheet_keeps_existing_workbook_behavior(self):
        content = workbook_bytes({"Empty": [], "csv": [], "Data": [["value"], [1]]})
        sheets = load_workbook(content, "empty.xlsx")
        self.assertTrue(sheets["Empty"].empty)
        self.assertTrue(sheets["csv"].empty)
        with self.assertRaisesRegex(ValueError, "header_row.*outside"):
            load_sheet(content, "empty.xlsx", "Empty", header_row=1)

    def test_excel_formulas_are_not_evaluated_and_macro_loading_is_disabled(self):
        content = workbook_bytes({"Sheet": [["ID", "calculated"], ["001", "=1+2"]]})
        for filename in ("formula.XLSX", "formula.xlsm"):
            with self.subTest(filename=filename):
                with patch.object(openpyxl, "load_workbook", wraps=openpyxl.load_workbook) as reader:
                    result = load_sheet(content, filename, "Sheet")
                self.assertTrue(reader.call_args.kwargs["data_only"])
                self.assertTrue(reader.call_args.kwargs["read_only"])
                self.assertFalse(reader.call_args.kwargs["keep_links"])
                self.assertFalse(reader.call_args.kwargs["keep_vba"])
                self.assertEqual(result["ID"].iloc[0], "001")
                value = result["calculated"].iloc[0]
                self.assertTrue(pd.isna(value) or value == "")
        csv = clean_frame(load_sheet(b"expression\n=1+2\n", "formula.csv", "csv"))
        self.assertEqual(csv["expression"].iloc[0], "=1+2")


if __name__ == "__main__":
    unittest.main()